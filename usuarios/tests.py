from datetime import date, datetime, timedelta
from io import BytesIO
from decimal import Decimal

from openpyxl import load_workbook
from django.test import Client, RequestFactory, TestCase
from django.urls import reverse
from django.utils import timezone

from inventario.models import Insumo, MovimientoInventario
from carta.models import Categoria, Plato
from pedidos.models import Factura, Mesa, Pedido
from usuarios.models import ProductoIPVTurno, Usuario, Notificacion, Turno, configurar_turno, get_turno_abierto, registrar_inicio_turno, cerrar_turno, mesas_del_turno
from usuarios.views import historial_turnos_view, historial_turno_detalle_view, ipv_turno


class UsuarioTurnoTests(TestCase):
    def setUp(self):
        self.admin = Usuario.objects.create_user(
            username='admin', password='123456', role='administrador'
        )
        self.mesera = Usuario.objects.create_user(
            username='mesera1', password='123456', role='mesera'
        )
        self.cocina = Usuario.objects.create_user(
            username='cocina1', password='123456', role='cocina'
        )
        self.cajera = Usuario.objects.create_user(
            username='cajera1', password='123456', role='cajera'
        )

    def test_registrar_inicio_turno_agrega_usuario_al_turno_actual(self):
        turno = registrar_inicio_turno(self.mesera)

        self.assertIsNotNone(turno)
        self.assertEqual(turno.usuarios.count(), 1)
        self.assertIn(self.mesera, turno.usuarios.all())

    def test_administrador_abre_el_turno_del_dia_expresamente(self):
        self.client.force_login(self.admin)

        response = self.client.post(reverse('usuarios:abrir_turno'))

        self.assertRedirects(response, reverse('usuarios:dashboard'))
        turno = Turno.objects.get(fecha=timezone.localdate())
        self.assertEqual(turno.estado, 'abierto')
        self.assertIn(self.admin, turno.usuarios.all())

    def test_iniciar_sesion_no_abre_un_turno_automaticamente(self):
        response = self.client.post(reverse('usuarios:login'), {
            'username': self.mesera.username,
            'password': '123456',
        })

        self.assertRedirects(response, reverse('pedidos:mesera'))
        self.assertFalse(Turno.objects.filter(fecha=timezone.localdate()).exists())

    def test_dashboard_muestra_facturas_del_turno_y_excluye_las_anteriores(self):
        turno = registrar_inicio_turno(self.mesera)
        mesa = Mesa.objects.create(numero=1, abierta=True)
        pedido_anterior = Pedido.objects.create(mesa=mesa, estado='cerrado')
        factura_anterior = Factura.objects.create(
            pedido=pedido_anterior,
            mesa_numero=mesa.numero,
            forma_pago='efectivo_cup',
            total_cup=Decimal('12.00'),
        )
        Factura.objects.filter(pk=factura_anterior.pk).update(
            creado_en=turno.apertura - timedelta(seconds=1)
        )
        pedido_turno = Pedido.objects.create(mesa=mesa, estado='cerrado')
        categoria = Categoria.objects.create(nombre='Bebidas IPV')
        plato_ipv = Plato.objects.create(
            categoria=categoria,
            nombre='Plato de prueba',
            precio=Decimal('25.00'),
        )
        factura_turno = Factura.objects.create(
            pedido=pedido_turno,
            mesa_numero=mesa.numero,
            forma_pago='efectivo_cup',
            total_cup=Decimal('25.00'),
            monto_efectivo_cup=Decimal('25.00'),
            cajera_nombre='Cajera de prueba',
            items_snapshot=[{
                'plato_id': plato_ipv.pk,
                'plato': 'Plato de prueba',
                'cantidad': 1,
                'precio_unit': '25.00',
                'subtotal': '25.00',
            }],
        )

        self.client.force_login(self.admin)
        response = self.client.get(reverse('usuarios:dashboard'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'name="fondo_inicial_cup"')
        self.assertContains(response, 'Detalle del turno actual')
        self.assertContains(response, f'href="{reverse("inventario:administrador")}#wifi"')
        self.assertEqual(list(response.context['facturas_turno']), [factura_turno])
        self.assertContains(response, f'Factura #{factura_turno.pk}')
        self.assertContains(response, 'Plato de prueba')
        self.assertContains(response, 'Inventario perpetuo de ventas')
        self.assertContains(response, 'Fondo inicial')
        self.assertContains(response, 'Saldo inicial')
        self.assertNotContains(response, 'Configurar IPV del turno')
        self.assertContains(response, 'setInterval(actualizarInventarioTurno, 5000)')
        self.assertNotContains(response, f'Factura #{factura_anterior.pk}')

        self.client.force_login(self.cajera)
        response_cajera = self.client.get(reverse('usuarios:dashboard'))
        self.assertEqual(response_cajera.status_code, 200)
        self.assertTrue(response_cajera.context['puede_ver_datos_financieros'])
        self.assertContains(response_cajera, f'Factura #{factura_turno.pk}')
        self.assertContains(response_cajera, 'Inventario perpetuo de ventas')

    def test_mesera_y_cocina_no_ven_datos_financieros_ni_ipv(self):
        turno = registrar_inicio_turno(self.mesera)
        turno.usuarios.add(self.cocina)
        turno.resumen = 'Total general: 987.65 CUP'
        turno.save(update_fields=['resumen'])
        mesa = Mesa.objects.create(numero=51, abierta=True)
        factura = Factura.objects.create(
            pedido=Pedido.objects.create(mesa=mesa, estado='cerrado'),
            mesa_numero=mesa.numero,
            forma_pago='efectivo_cup',
            total_cup=Decimal('987.65'),
            monto_efectivo_cup=Decimal('987.65'),
            items_snapshot=[{
                'plato': 'Venta privada',
                'cantidad': 1,
                'precio_unit': '987.65',
                'subtotal': '987.65',
            }],
        )

        for usuario in (self.mesera, self.cocina):
            with self.subTest(rol=usuario.role):
                self.client.force_login(usuario)
                response = self.client.get(reverse('usuarios:dashboard'))

                self.assertEqual(response.status_code, 200)
                self.assertFalse(response.context['puede_ver_datos_financieros'])
                self.assertFalse(response.context['facturas_turno'].exists())
                self.assertNotContains(response, 'Ventas totales')
                self.assertNotContains(response, 'Facturas del turno')
                self.assertNotContains(response, 'Inventario perpetuo de ventas')
                self.assertNotContains(response, '987.65')
                self.assertNotContains(response, f'Factura #{factura.pk}')
                self.assertEqual(
                    self.client.get(reverse('usuarios:inventario_ventas_turno')).status_code,
                    403,
                )
                self.assertEqual(
                    self.client.get(reverse('usuarios:ipv_turno', args=[turno.pk])).status_code,
                    403,
                )
                self.assertEqual(
                    self.client.get(reverse('usuarios:historial_turno_detalle', args=[turno.pk])).status_code,
                    302,
                )
                self.assertEqual(
                    self.client.get(reverse('usuarios:historial_turnos')).status_code,
                    302,
                )

    def test_cajera_no_puede_ver_turnos_cerrados_anteriores(self):
        turno_cerrado = Turno.objects.create(
            fecha=timezone.localdate() - timedelta(days=1),
            estado='cerrado',
            cierre=timezone.now(),
        )
        turno_actual = registrar_inicio_turno(self.cajera)
        self.client.force_login(self.cajera)

        dashboard = self.client.get(reverse('usuarios:dashboard'))

        self.assertEqual(dashboard.status_code, 200)
        self.assertNotContains(dashboard, reverse('usuarios:historial_turnos'))
        self.assertContains(
            dashboard,
            reverse('usuarios:historial_turno_detalle', args=[turno_actual.pk]),
        )
        self.assertEqual(self.client.get(reverse('usuarios:historial_turnos')).status_code, 302)
        self.assertEqual(
            self.client.get(reverse('usuarios:historial_turno_detalle', args=[turno_cerrado.pk])).status_code,
            302,
        )
        for nombre_ruta in ('ipv_turno', 'ipv_turno_pdf', 'ipv_turno_excel'):
            with self.subTest(ruta=nombre_ruta):
                self.assertEqual(
                    self.client.get(reverse(f'usuarios:{nombre_ruta}', args=[turno_cerrado.pk])).status_code,
                    403,
                )
        self.assertEqual(
            self.client.get(reverse('usuarios:historial_turno_detalle', args=[turno_actual.pk])).status_code,
            200,
        )

    def test_historial_de_turnos_se_pagina(self):
        for indice in range(11):
            turno = Turno.objects.create(estado='cerrado')
            Turno.objects.filter(pk=turno.pk).update(
                cierre=timezone.now() + timedelta(seconds=indice),
            )

        self.client.force_login(self.admin)
        response = self.client.get(reverse('usuarios:historial_turnos'), {'page': 2})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['historial'].number, 2)
        self.assertEqual(response.context['historial'].paginator.count, 11)
        self.assertEqual(len(response.context['historial']), 1)

    def test_resumen_acumulado_respeta_el_rango_de_fechas(self):
        hoy = timezone.localdate()
        turnos = [
            (hoy - timedelta(days=3), 'efectivo_cup', Decimal('100.00')),
            (hoy - timedelta(days=2), 'transferencia', Decimal('200.00')),
            (hoy - timedelta(days=1), 'efectivo_cup', Decimal('300.00')),
        ]
        mesa = Mesa.objects.create(numero=72, abierta=True)

        for fecha_turno, forma_pago, total in turnos:
            turno = Turno.objects.create(estado='cerrado', fecha=fecha_turno)
            apertura = timezone.make_aware(datetime.combine(fecha_turno, datetime.min.time()))
            cierre = apertura + timedelta(days=1) - timedelta(microseconds=1)
            Turno.objects.filter(pk=turno.pk).update(apertura=apertura, cierre=cierre)
            factura = Factura.objects.create(
                pedido=Pedido.objects.create(mesa=mesa, estado='cerrado'),
                mesa_numero=mesa.numero,
                forma_pago=forma_pago,
                total_cup=total,
                monto_efectivo_cup=total if forma_pago == 'efectivo_cup' else Decimal('0'),
                monto_transferencia_cup=total if forma_pago == 'transferencia' else Decimal('0'),
            )
            creado_en = timezone.make_aware(datetime.combine(fecha_turno, datetime.min.time()))
            Factura.objects.filter(pk=factura.pk).update(creado_en=creado_en)

        self.client.force_login(self.admin)
        response = self.client.get(reverse('usuarios:historial_turnos'), {
            'fecha_desde': (hoy - timedelta(days=3)).isoformat(),
            'fecha_hasta': (hoy - timedelta(days=2)).isoformat(),
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['historial'].paginator.count, 2)
        self.assertEqual(response.context['resumen_general']['turnos_cerrados'], 2)
        self.assertEqual(response.context['resumen_general']['total_general_cup'], Decimal('300.00'))
        self.assertEqual(response.context['resumen_general']['total_efectivo_cup'], Decimal('100.00'))
        self.assertEqual(response.context['resumen_general']['total_transferencia_cup'], Decimal('200.00'))
        self.assertContains(response, (hoy - timedelta(days=3)).isoformat())
        self.assertContains(response, (hoy - timedelta(days=2)).isoformat())

    def test_reportes_muestran_ventas_y_categorias_del_periodo_seleccionado(self):
        hoy = timezone.localdate()
        categoria_bebidas = Categoria.objects.create(nombre='Bebidas')
        categoria_comidas = Categoria.objects.create(nombre='Comidas')
        plato = Plato.objects.create(nombre='Jugo', categoria=categoria_bebidas, precio='25.00')
        otro_plato = Plato.objects.create(nombre='Sándwich', categoria=categoria_comidas, precio='50.00')
        mesa = Mesa.objects.create(numero=73, abierta=True)
        ventas = [
            (
                hoy - timedelta(days=1),
                'efectivo_cup',
                Decimal('100.00'),
                Decimal('100.00'),
                Decimal('0'),
                Decimal('0'),
                [
                    {'plato': plato.nombre, 'categoria': categoria_bebidas.nombre, 'cantidad': 2},
                    {'plato': otro_plato.nombre, 'categoria': categoria_comidas.nombre, 'cantidad': 1},
                ],
            ),
            (
                hoy,
                'transferencia',
                Decimal('200.00'),
                Decimal('0'),
                Decimal('200.00'),
                Decimal('0'),
                [{'plato': plato.nombre, 'categoria': categoria_bebidas.nombre, 'cantidad': 3}],
            ),
            (
                hoy - timedelta(days=8),
                'usd',
                Decimal('75.00'),
                Decimal('0'),
                Decimal('0'),
                Decimal('3.00'),
                [{'plato': plato.nombre, 'categoria': categoria_bebidas.nombre, 'cantidad': 4}],
            ),
        ]
        for fecha, forma_pago, total, efectivo, transferencia, usd, items in ventas:
            factura = Factura.objects.create(
                pedido=Pedido.objects.create(mesa=mesa, estado='cerrado'),
                mesa_numero=mesa.numero,
                forma_pago=forma_pago,
                total_cup=total,
                monto_efectivo_cup=efectivo,
                monto_transferencia_cup=transferencia,
                monto_usd=usd,
                items_snapshot=items,
            )
            creado_en = timezone.make_aware(datetime.combine(fecha, datetime.min.time()))
            Factura.objects.filter(pk=factura.pk).update(creado_en=creado_en)

        self.client.force_login(self.admin)
        response = self.client.get(reverse('usuarios:reportes'), {'periodo': 'semana'})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, f'href="{reverse("usuarios:reportes")}"')
        self.assertContains(response, f'href="{reverse("inventario:administrador")}#wifi"')
        self.assertEqual(response.context['cantidad_facturas'], 2)
        self.assertEqual(response.context['total_cup'], Decimal('300.00'))
        self.assertEqual(response.context['total_unidades'], 6)
        self.assertEqual(response.context['productos_mas_vendidos'][0]['cantidad'], 5)
        self.assertEqual(response.context['ventas_por_categoria'][0]['participacion'], 83)
        self.assertIn('#34d399', response.context['categorias_donut'])
        self.assertTrue(response.context['pagos_donut'])
        categorias = {
            categoria['nombre']: categoria['cantidad']
            for categoria in response.context['ventas_por_categoria']
        }
        self.assertEqual(categorias, {'Bebidas': 5, 'Comidas': 1})
        self.assertContains(response, 'Semana actual')
        self.assertContains(response, 'Recaudación por día')
        self.assertContains(response, 'report-timeline-bar')
        self.assertContains(response, 'report-donut')

        response_rango = self.client.get(reverse('usuarios:reportes'), {
            'periodo': 'rango',
            'fecha_desde': hoy.isoformat(),
            'fecha_hasta': hoy.isoformat(),
        })
        self.assertEqual(response_rango.context['cantidad_facturas'], 1)
        self.assertEqual(response_rango.context['total_cup'], Decimal('200.00'))

    def test_reportes_son_solo_para_el_administrador_y_validan_el_rango(self):
        self.client.force_login(self.cajera)
        self.assertEqual(self.client.get(reverse('usuarios:reportes')).status_code, 302)

        self.client.force_login(self.admin)
        response = self.client.get(reverse('usuarios:reportes'), {
            'periodo': 'rango',
            'fecha_desde': '2026-12-31',
            'fecha_hasta': '2026-01-01',
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'La fecha inicial debe ser anterior o igual a la fecha final.')

    def test_dashboard_de_mesera_usa_navegacion_movil_del_panel(self):
        self.client.force_login(self.mesera)

        response = self.client.get(reverse('usuarios:dashboard'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'class="mesera-app bg-slate-900')
        self.assertContains(response, 'turn-mobile-nav')
        self.assertContains(response, 'aria-label="Cerrar sesión"')
        self.assertNotContains(response, 'responsive-sidebar.js')

    def test_dashboard_admin_sin_turno_muestra_fecha_apertura_y_turno_anterior(self):
        turno_anterior = Turno.objects.create(
            fecha=timezone.localdate() - timedelta(days=1),
            estado='cerrado',
            cierre=timezone.now(),
            resumen='Cierre anterior de prueba',
        )
        self.client.force_login(self.admin)

        response = self.client.get(reverse('usuarios:dashboard'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Abrir turno de hoy')
        self.assertContains(response, 'El turno anterior fue el')
        self.assertContains(response, reverse('usuarios:historial_turno_detalle', args=[turno_anterior.pk]))
        self.assertContains(response, f'href="{reverse("usuarios:reportes")}"')
        self.assertContains(response, f'href="{reverse("inventario:administrador")}#inventario-restaurante"')
        self.assertContains(response, f'href="{reverse("inventario:administrador")}#inventario-cocina"')
        self.assertNotContains(response, 'Crear usuario')

    def test_dashboard_admin_con_turno_cerrado_muestra_detalle_e_historial(self):
        turno = Turno.objects.create(
            fecha=timezone.localdate(),
            estado='cerrado',
            cierre=timezone.now(),
        )
        self.client.force_login(self.admin)

        response = self.client.get(reverse('usuarios:dashboard'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Cerrado')
        self.assertContains(response, reverse('usuarios:historial_turno_detalle', args=[turno.pk]))
        self.assertContains(response, reverse('usuarios:historial_turnos'))
        self.assertNotContains(response, 'Abrir turno de hoy')

    def test_historial_turnos_filtra_desde_el_servidor_y_mantiene_los_filtros(self):
        turno_mayo_10 = Turno.objects.create(
            fecha=date(2024, 5, 10),
            estado='cerrado',
            cierre=timezone.now(),
            resumen='Inventario temprano',
        )
        turno_mayo_20 = Turno.objects.create(
            fecha=date(2024, 5, 20),
            estado='cerrado',
            cierre=timezone.now(),
            resumen='Cierre nocturno',
        )
        turno_junio = Turno.objects.create(
            fecha=date(2024, 6, 10),
            estado='cerrado',
            cierre=timezone.now(),
            resumen='Inventario de junio',
        )
        turno_mayo_10.usuarios.add(self.mesera)
        turno_mayo_20.usuarios.add(self.cajera)
        turno_junio.usuarios.add(self.mesera)
        self.client.force_login(self.admin)

        response = self.client.get(reverse('usuarios:historial_turnos'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Apertura:')
        self.assertContains(response, 'Cierre:')
        self.assertContains(
            response,
            f"Apertura: {timezone.localtime(turno_mayo_10.apertura):%d/%m/%Y %H:%M}",
        )
        self.assertContains(
            response,
            f"Cierre: {timezone.localtime(turno_mayo_10.cierre):%d/%m/%Y %H:%M}",
        )
        self.assertEqual(response.content.count(b'<input '), 3)
        self.assertEqual(response.content.count(b'<select '), 2)
        self.assertEqual(response.content.count(b'<button '), 2)
        self.assertEqual(
            {item['turno'].pk for item in response.context['historial']},
            {turno_mayo_10.pk, turno_mayo_20.pk, turno_junio.pk},
        )
        self.assertContains(response, 'id="filtro-texto"')
        self.assertContains(response, 'id="filtro-mes"')
        self.assertContains(response, 'id="filtro-anio"')
        self.assertContains(response, 'id="filtro-fecha-desde"')
        self.assertContains(response, 'id="filtro-fecha-hasta"')
        self.assertContains(response, 'Filtrar por rango de fechas')
        self.assertContains(response, 'aria-label="Fecha desde"')
        self.assertContains(response, 'aria-label="Fecha hasta"')
        self.assertContains(response, 'aria-label="Mes"')
        self.assertContains(response, 'aria-label="Año"')
        self.assertContains(response, 'background-color: #0f172a;')
        self.assertContains(response, 'color: #f8fafc;')
        self.assertContains(response, 'Filtrar por mes y año')
        self.assertContains(response, 'id="turn-filter-form"')
        self.assertContains(response, 'id="limpiar-filtros"')
        self.assertContains(response, 'method="get"')
        self.assertContains(response, 'name="fecha_desde"')
        self.assertContains(response, 'name="fecha_hasta"')
        self.assertContains(response, "addEventListener('click'")

        response_filtrada = self.client.get(reverse('usuarios:historial_turnos'), {
            'fecha_desde': '2024-05-10',
            'fecha_hasta': '2024-05-20',
        })
        self.assertEqual(response_filtrada.status_code, 200)
        self.assertEqual(response_filtrada.context['historial'].paginator.count, 2)
        self.assertEqual(response_filtrada.context['resumen_general']['turnos_cerrados'], 2)
        self.assertEqual(
            {item['turno'].pk for item in response_filtrada.context['historial']},
            {turno_mayo_10.pk, turno_mayo_20.pk},
        )
        self.assertContains(response_filtrada, 'value="2024-05-10"')
        self.assertContains(response_filtrada, 'value="2024-05-20"')

    def test_admin_tiene_seccion_separada_para_gestionar_usuarios(self):
        self.client.force_login(self.admin)

        response = self.client.get(reverse('usuarios:gestion_usuarios'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Crear usuario')
        self.assertContains(response, 'Usuarios registrados')

    def test_editores_de_usuario_solo_permiten_una_modal_abierta(self):
        self.client.force_login(self.admin)

        response = self.client.get(reverse('usuarios:gestion_usuarios'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '<details class="turn-user-edit" name="user-edit">')
        self.assertContains(response, "otherDialog.open = false")

    def test_sidebar_del_administrador_se_mantiene_completo_en_sus_secciones(self):
        turno = Turno.objects.create(
            fecha=timezone.localdate() - timedelta(days=1),
            estado='cerrado',
            cierre=timezone.now(),
        )
        self.client.force_login(self.admin)
        paginas = [
            reverse('usuarios:gestion_usuarios'),
            reverse('usuarios:historial_turnos'),
            reverse('usuarios:historial_turno_detalle', args=[turno.pk]),
        ]
        enlaces_sidebar = [
            f'href="{reverse("inventario:administrador")}#configuracion"',
            f'href="{reverse("inventario:administrador")}#carta"',
            f'href="{reverse("inventario:administrador")}#inventario-restaurante"',
            f'href="{reverse("inventario:administrador")}#inventario-cocina"',
            f'href="{reverse("usuarios:gestion_usuarios")}"',
            f'href="{reverse("usuarios:dashboard")}"',
            f'href="{reverse("usuarios:historial_turnos")}"',
        ]

        for pagina in paginas:
            with self.subTest(pagina=pagina):
                response = self.client.get(pagina)

                self.assertEqual(response.status_code, 200)
                for enlace in enlaces_sidebar:
                    self.assertContains(response, enlace)
                if pagina == reverse('usuarios:gestion_usuarios'):
                    self.assertContains(response, 'admin')
                    self.assertContains(response, 'Administrador &bull; Activo')

    def test_detalle_diario_muestra_navegacion_segun_el_rol(self):
        roles = (
            (self.mesera, 'Mesera', reverse('pedidos:mesera')),
            (self.cocina, 'Cocina', reverse('pedidos:cocina')),
            (self.admin, 'Administrador', reverse('inventario:administrador')),
        )

        for usuario, etiqueta_rol, enlace_rol in roles:
            with self.subTest(rol=usuario.role):
                self.client.force_login(usuario)
                response = self.client.get(reverse('usuarios:dashboard'))

                self.assertEqual(response.status_code, 200)
                self.assertContains(response, enlace_rol)
                self.assertContains(response, f'{etiqueta_rol} &bull; Activo')
                if usuario.role == 'cocina':
                    self.assertContains(
                        response,
                        f'href="{reverse("pedidos:cocina")}?vista=ventas"',
                    )
                if usuario.role == 'mesera':
                    for seccion in (
                        'tables',
                        'exterior-tables',
                        'active-orders',
                        'bar',
                        'takeaway',
                        'menu',
                    ):
                        self.assertContains(
                            response,
                            f'href="{reverse("pedidos:mesera")}#{seccion}"',
                        )
                if usuario.role != 'administrador':
                    self.assertNotContains(response, f'href="{reverse("inventario:administrador")}#configuracion"')
                    self.assertNotContains(response, f'href="{reverse("usuarios:gestion_usuarios")}"')
                    self.assertNotContains(response, f'href="{reverse("usuarios:historial_turnos")}"')

    def test_reporte_inventario_turno_calcula_ventas_entradas_roturas_y_saldos(self):
        categoria = Categoria.objects.create(nombre='Bebidas IPV', orden=1)
        plato = Plato.objects.create(categoria=categoria, nombre='Coca-Cola', precio=Decimal('25.00'))
        insumo = Insumo.objects.create(nombre='Coca-Cola', stock_actual=Decimal('10'))
        cierre_anterior = timezone.now() - timedelta(days=1)
        turno_anterior = Turno.objects.create(
            fecha=timezone.localdate() - timedelta(days=1),
            estado='cerrado',
            cierre=cierre_anterior,
        )
        Turno.objects.filter(pk=turno_anterior.pk).update(
            apertura=cierre_anterior - timedelta(hours=8),
        )
        ProductoIPVTurno.objects.create(
            turno=turno_anterior,
            plato=plato,
            cantidad_inicial=Decimal('12'),
            entradas=Decimal('4'),
            merma_rotura=Decimal('1'),
        )
        pedido_anterior = Pedido.objects.create(mesa=Mesa.objects.create(numero=50), estado='cerrado')
        factura_anterior = Factura.objects.create(
            pedido=pedido_anterior,
            mesa_numero=pedido_anterior.mesa.numero,
            forma_pago='efectivo_cup',
            total_cup=Decimal('75.00'),
            items_snapshot=[{
                'plato_id': plato.pk,
                'plato': plato.nombre,
                'cantidad': 3,
                'precio_unit': '25.00',
                'subtotal': '75.00',
            }],
        )
        Factura.objects.filter(pk=factura_anterior.pk).update(
            creado_en=cierre_anterior - timedelta(minutes=1),
        )
        turno = registrar_inicio_turno(self.cajera)
        self.client.force_login(self.admin)

        dashboard = self.client.get(reverse('usuarios:dashboard'))

        self.assertEqual(dashboard.status_code, 200)
        self.assertEqual(dashboard.context['opciones_ipv'][0]['cantidad_inicial'], Decimal('12'))

        self.client.force_login(self.cajera)
        configuracion = self.client.post(reverse('usuarios:configurar_ipv_turno'), {
            'fondo_inicial_cup': '100.00',
            'platos': [str(plato.pk)],
        }, follow=True)

        self.assertContains(configuracion, 'El IPV del turno quedó configurado.')
        registro = ProductoIPVTurno.objects.get(turno=turno, plato=plato)
        self.assertEqual(turno.__class__.objects.get(pk=turno.pk).fondo_inicial_cup, Decimal('100.00'))
        self.assertEqual(registro.cantidad_inicial, Decimal('12'))
        self.assertEqual(registro.entradas, Decimal('0'))

        mesa = Mesa.objects.create(numero=51, abierta=True)
        pedido = Pedido.objects.create(mesa=mesa, estado='cerrado')
        Factura.objects.create(
            pedido=pedido,
            mesa_numero=mesa.numero,
            forma_pago='mixto',
            total_cup=Decimal('75.00'),
            monto_efectivo_cup=Decimal('40.00'),
            monto_transferencia_cup=Decimal('35.00'),
            items_snapshot=[{
                'plato_id': plato.pk,
                'plato': plato.nombre,
                'cantidad': 3,
                'precio_unit': '25.00',
                'subtotal': '75.00',
            }],
        )
        respuesta_merma = self.client.post(
            reverse('usuarios:registrar_merma_ipv', args=[registro.pk]),
            {'cantidad': '1'},
        )
        self.assertEqual(respuesta_merma.status_code, 200)
        respuesta_merma_excesiva = self.client.post(
            reverse('usuarios:registrar_merma_ipv', args=[registro.pk]),
            {'cantidad': '11'},
        )
        self.assertEqual(respuesta_merma_excesiva.status_code, 400)
        registro.refresh_from_db()
        self.assertEqual(registro.merma_rotura, Decimal('1'))

        response = self.client.get(reverse('usuarios:inventario_ventas_turno'))

        self.assertEqual(response.status_code, 200)
        datos = response.json()
        self.assertEqual([row['nombre'] for row in datos['productos']], [plato.nombre])
        producto = datos['productos'][0]
        self.assertEqual(producto['saldo_inicial'], '12.000')
        self.assertEqual(producto['entradas'], '0.000')
        self.assertEqual(producto['cantidad_vendida'], '3')
        self.assertEqual(producto['merma_rotura'], '1.000')
        self.assertEqual(producto['saldo_final'], '8.000')
        self.assertEqual(producto['precio_venta'], '25.00')
        self.assertEqual(producto['importe_venta'], '75.00')
        self.assertEqual(datos['fondo_inicial_cup'], '100.00')
        self.assertEqual(datos['efectivo_final_cup'], '140.00')
        self.assertEqual(datos['ventas_transferencia_cup'], '35.00')
        self.assertEqual(insumo.stock_actual, Decimal('10'))
        self.assertFalse(MovimientoInventario.objects.exists())

    def test_ipv_se_configura_con_saldo_final_del_turno_anterior_y_sin_entradas_manuales(self):
        categoria = Categoria.objects.create(nombre='Platos IPV', orden=1)
        plato = Plato.objects.create(categoria=categoria, nombre='Ropa vieja', precio=Decimal('18.00'))
        cierre_anterior = timezone.now() - timedelta(days=1)
        turno_anterior = Turno.objects.create(
            fecha=timezone.localdate() - timedelta(days=1),
            estado='cerrado',
            cierre=cierre_anterior,
        )
        Turno.objects.filter(pk=turno_anterior.pk).update(
            apertura=cierre_anterior - timedelta(hours=8),
        )
        ProductoIPVTurno.objects.create(
            turno=turno_anterior,
            plato=plato,
            cantidad_inicial=Decimal('10'),
            entradas=Decimal('0'),
            merma_rotura=Decimal('2'),
        )
        pedido_anterior = Pedido.objects.create(mesa=Mesa.objects.create(numero=90), estado='cerrado')
        Factura.objects.create(
            pedido=pedido_anterior,
            mesa_numero=pedido_anterior.mesa.numero,
            forma_pago='efectivo_cup',
            total_cup=Decimal('54.00'),
            items_snapshot=[{
                'plato_id': plato.pk,
                'plato': plato.nombre,
                'cantidad': 3,
                'precio_unit': '18.00',
                'subtotal': '54.00',
            }],
        )
        Factura.objects.filter(pk__in=Factura.objects.filter(pedido=pedido_anterior).values_list('pk', flat=True)).update(
            creado_en=cierre_anterior - timedelta(minutes=1),
        )

        turno = registrar_inicio_turno(self.cajera)
        self.client.force_login(self.admin)
        dashboard = self.client.get(reverse('usuarios:dashboard'))
        self.assertEqual(dashboard.context['opciones_ipv'][0]['cantidad_inicial'], Decimal('5'))

        self.client.force_login(self.cajera)
        response = self.client.post(reverse('usuarios:configurar_ipv_turno'), {
            'fondo_inicial_cup': '150.00',
            'platos': [str(plato.pk)],
            f'inicial_{plato.pk}': '999',
            f'entradas_{plato.pk}': '77',
        }, follow=True)

        self.assertContains(response, 'El IPV del turno quedó configurado.')
        registro = ProductoIPVTurno.objects.get(turno=turno, plato=plato)
        self.assertEqual(registro.cantidad_inicial, Decimal('5'))
        self.assertEqual(registro.entradas, Decimal('0'))

    def test_cerrar_turno_crea_resumen_y_notifica_al_admin(self):
        turno = registrar_inicio_turno(self.mesera)
        registrar_inicio_turno(self.cajera)
        turno.fondo_inicial_cup = Decimal('75.00')
        turno.save(update_fields=['fondo_inicial_cup'])

        mesa = Mesa.objects.create(numero=10, abierta=True)
        pedido = Pedido.objects.create(mesa=mesa, estado='cerrado')
        Pedido.objects.create(mesa=mesa, estado='cancelado')
        Pedido.objects.create(modalidad='para_llevar', estado='cancelado')
        Factura.objects.create(
            pedido=pedido,
            mesa_numero=mesa.numero,
            forma_pago='efectivo_cup',
            total_cup=Decimal('250.00'),
            monto_efectivo_cup=Decimal('250.00'),
        )

        turno_cerrado = cerrar_turno(self.cajera)

        self.assertEqual(turno_cerrado.estado, 'cerrado')
        self.assertIn('250.00', turno_cerrado.resumen)
        self.assertIn('Fondo inicial: 75.00 CUP', turno_cerrado.resumen)
        self.assertIn('Efectivo esperado: 325.00 CUP', turno_cerrado.resumen)
        self.assertTrue(Notificacion.objects.filter(
            destinatario=self.admin,
            turno=turno,
            asunto='Cierre de turno realizado',
            mensaje__contains=f'Cajera {self.cajera.username} cerró el turno',
        ).exists())

    def test_resumen_financiero_incluye_facturas_de_turno_que_cruza_medianoche(self):
        turno = Turno.objects.create(estado='cerrado')
        apertura = timezone.now().replace(hour=23, minute=30, second=0, microsecond=0) - timedelta(days=1)
        cierre = apertura + timedelta(hours=3)
        Turno.objects.filter(pk=turno.pk).update(apertura=apertura, cierre=cierre)
        turno.refresh_from_db()
        mesa = Mesa.objects.create(numero=12)
        factura_antes = Factura.objects.create(
            pedido=Pedido.objects.create(mesa=mesa, estado='cerrado'),
            mesa_numero=mesa.numero,
            forma_pago='efectivo_cup',
            total_cup=Decimal('99.00'),
            monto_efectivo_cup=Decimal('99.00'),
        )
        factura_despues_de_medianoche = Factura.objects.create(
            pedido=Pedido.objects.create(mesa=mesa, estado='cerrado'),
            mesa_numero=mesa.numero,
            forma_pago='efectivo_cup',
            total_cup=Decimal('25.00'),
            monto_efectivo_cup=Decimal('25.00'),
        )
        Factura.objects.filter(pk=factura_antes.pk).update(creado_en=apertura - timedelta(minutes=1))
        Factura.objects.filter(pk=factura_despues_de_medianoche.pk).update(
            creado_en=apertura + timedelta(hours=2),
        )

        self.assertEqual(turno.total_general_cup, Decimal('25.00'))
        self.assertEqual(turno.total_efectivo_cup, Decimal('25.00'))
        self.assertEqual(turno.cantidad_facturas, 1)

    def test_cerrar_turno_rechaza_pedidos_pendientes_en_mesa(self):
        turno = registrar_inicio_turno(self.cajera)
        mesa = Mesa.objects.create(numero=10, abierta=True)
        Pedido.objects.create(mesa=mesa, estado='servido')

        with self.assertRaisesRegex(ValueError, 'pedidos pendientes en una mesa'):
            cerrar_turno(self.cajera)

        turno.refresh_from_db()
        self.assertEqual(turno.estado, 'abierto')

    def test_cerrar_turno_rechaza_ticket_para_llevar_sin_cobrar(self):
        turno = registrar_inicio_turno(self.cajera)
        Pedido.objects.create(modalidad='para_llevar', estado='entregado')

        with self.assertRaisesRegex(ValueError, 'tickets para llevar sin cobrar'):
            cerrar_turno(self.cajera)

        turno.refresh_from_db()
        self.assertEqual(turno.estado, 'abierto')

    def test_administrador_puede_cerrar_turno_desde_el_panel(self):
        turno = registrar_inicio_turno(self.admin)
        self.client.force_login(self.admin)

        dashboard = self.client.get(reverse('usuarios:dashboard'))
        self.assertContains(dashboard, 'Cerrar turno')

        response = self.client.post(reverse('usuarios:cerrar_turno'), {
            'observaciones': 'Cierre realizado por administración',
        }, follow=True)

        self.assertRedirects(
            response,
            reverse('usuarios:historial_turno_detalle', args=[turno.pk]),
        )
        self.assertContains(response, 'Turno cerrado')
        self.assertContains(response, 'Cierre realizado por administración')
        turno.refresh_from_db()
        self.assertEqual(turno.estado, 'cerrado')
        self.assertEqual(turno.observaciones, 'Cierre realizado por administración')
        self.assertTrue(Notificacion.objects.filter(
            destinatario=self.admin,
            mensaje__contains='Administrador admin cerró el turno',
        ).exists())

    def test_administrador_puede_cerrar_turno_abierto_del_dia_anterior(self):
        turno_anterior = Turno.objects.create(
            fecha=timezone.localdate() - timedelta(days=1),
            estado='abierto',
        )
        turno_hoy = Turno.objects.create(fecha=timezone.localdate(), estado='abierto')
        self.client.force_login(self.admin)

        dashboard = self.client.get(reverse('usuarios:dashboard'))

        self.assertContains(dashboard, 'El turno pasó de medianoche')
        self.assertContains(dashboard, 'Continuar turno')
        self.assertContains(dashboard, 'Cerrar turno')
        self.assertContains(dashboard, f'name="turno_id" value="{turno_anterior.pk}"')

        response = self.client.post(reverse('usuarios:cerrar_turno'), {
            'turno_id': turno_anterior.pk,
            'observaciones': 'Cierre del turno atrasado',
        }, follow=True)

        self.assertRedirects(
            response,
            reverse('usuarios:historial_turno_detalle', args=[turno_anterior.pk]),
        )
        self.assertContains(response, 'Turno cerrado')
        self.assertContains(response, 'Cierre del turno atrasado')
        turno_anterior.refresh_from_db()
        turno_hoy.refresh_from_db()
        self.assertEqual(turno_anterior.estado, 'cerrado')
        self.assertEqual(turno_anterior.observaciones, 'Cierre del turno atrasado')
        self.assertEqual(turno_hoy.estado, 'abierto')
        self.assertTrue(Notificacion.objects.filter(
            destinatario=self.admin,
            turno=turno_anterior,
            mensaje__contains='Administrador admin cerró el turno',
        ).exists())

    def test_administrador_puede_confirmar_continuidad_de_turno_tras_medianoche(self):
        turno = Turno.objects.create(
            fecha=timezone.localdate() - timedelta(days=1),
            estado='abierto',
        )
        self.client.force_login(self.admin)

        response = self.client.get(reverse('usuarios:dashboard'))
        self.assertContains(response, 'El turno pasó de medianoche')
        self.assertNotContains(response, 'id="cierre-turno-heading"')

        response = self.client.post(reverse('usuarios:continuar_turno'), {
            'turno_id': turno.pk,
        }, follow=True)

        self.assertContains(response, 'El turno continuará abierto hasta que se realice su cierre')
        self.assertNotContains(response, 'El turno pasó de medianoche')
        turno.refresh_from_db()
        self.assertEqual(turno.estado, 'abierto')
        self.assertTrue(turno.continuidad_confirmada)

    def test_otros_roles_no_pueden_cerrar_turno(self):
        registrar_inicio_turno(self.mesera)
        self.client.force_login(self.mesera)

        response = self.client.post(
            reverse('usuarios:cerrar_turno'),
            follow=True,
        )

        self.assertContains(response, 'Solo caja o el administrador pueden cerrar el turno')
        self.assertEqual(Turno.objects.get(estado='abierto').estado, 'abierto')

    def test_configurar_turno_define_mesas_y_carta_del_turno(self):
        categoria = Categoria.objects.create(nombre='Entradas', orden=1)
        plato = Plato.objects.create(categoria=categoria, nombre='Croquetas', descripcion='Caseras', precio=Decimal('5.00'))

        turno = registrar_inicio_turno(self.admin)
        turno_configurado = configurar_turno(turno, cantidad_mesas=12, platos=[plato])

        self.assertEqual(turno_configurado.cantidad_mesas, 12)
        self.assertEqual(turno_configurado.platos.count(), 1)
        self.assertIn(plato, turno_configurado.platos.all())

    def test_cerrar_turno_incluye_resumen_detallado_de_ventas(self):
        turno = registrar_inicio_turno(self.mesera)
        registrar_inicio_turno(self.cajera)

        mesa = Mesa.objects.create(numero=10, abierta=True)
        pedido = Pedido.objects.create(mesa=mesa, estado='cerrado')
        Factura.objects.create(
            pedido=pedido,
            mesa_numero=mesa.numero,
            forma_pago='efectivo_cup',
            total_cup=Decimal('200.00'),
            monto_efectivo_cup=Decimal('200.00'),
        )
        Factura.objects.create(
            pedido=Pedido.objects.create(mesa=mesa, estado='cerrado'),
            mesa_numero=mesa.numero,
            forma_pago='transferencia',
            total_cup=Decimal('300.00'),
            monto_transferencia_cup=Decimal('300.00'),
        )
        Factura.objects.create(
            pedido=Pedido.objects.create(mesa=mesa, estado='cerrado'),
            mesa_numero=mesa.numero,
            forma_pago='usd',
            total_cup=Decimal('400.00'),
            monto_usd=Decimal('10.00'),
            tasa_cambio=Decimal('40.00'),
        )

        turno_cerrado = cerrar_turno(self.cajera)

        self.assertIn('USD', turno_cerrado.resumen)
        self.assertIn('Efectivo CUP', turno_cerrado.resumen)
        self.assertIn('Transferencia CUP', turno_cerrado.resumen)
        self.assertIn('Total general', turno_cerrado.resumen)
        self.assertIn('900.00', turno_cerrado.resumen)

    def test_mesas_del_turno_respeta_el_estado_individual_de_cada_mesa(self):
        for numero in range(1, 15):
            Mesa.objects.get_or_create(numero=numero, defaults={'activa': True})

        turno = registrar_inicio_turno(self.admin)
        configurar_turno(turno, cantidad_mesas=5, platos=[])
        Mesa.objects.filter(numero__in=[2, 4]).update(activa=False)

        mesas = mesas_del_turno(turno)

        self.assertEqual(
            list(mesas.values_list('numero', flat=True)),
            [1, 3, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14],
        )

    def test_mesera_ve_las_mesas_de_un_turno_abierto_de_un_dia_anterior(self):
        turno = Turno.objects.create(
            fecha=timezone.localdate() - timedelta(days=1),
            estado='abierto',
            cantidad_mesas=3,
        )
        self.client.force_login(self.mesera)

        response = self.client.get(reverse('pedidos:mesera'))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['turno'], turno)
        self.assertEqual(list(response.context['mesas'].values_list('numero_zona', flat=True)), [1, 2, 3])

    def test_mesas_activas_fuera_del_rango_anterior_siguen_disponibles(self):
        for numero in range(1, 15):
            Mesa.objects.get_or_create(numero=numero, defaults={'activa': True})

        turno = registrar_inicio_turno(self.admin)
        configurar_turno(turno, cantidad_mesas=5, platos=[])

        mesas = mesas_del_turno(turno)

        self.assertEqual(list(mesas.values_list('numero', flat=True)), list(range(1, 15)))

    def test_configurar_turno_crea_las_mesas_faltantes_hasta_la_cantidad(self):
        turno = registrar_inicio_turno(self.admin)
        configurar_turno(turno, cantidad_mesas=5, platos=[])

        mesas = mesas_del_turno(turno)

        self.assertEqual(list(mesas.values_list('numero', flat=True)), [1, 2, 3, 4, 5])
        self.assertEqual(Mesa.objects.filter(activa=True).count(), 5)

    def test_configurar_turno_crea_mesas_exteriores_separadas_y_etiquetadas(self):
        turno = registrar_inicio_turno(self.admin)
        configurar_turno(turno, cantidad_mesas=2, platos=[], cantidad_mesas_exteriores=3)

        mesas = mesas_del_turno(turno)
        mesas_salon = [mesa for mesa in mesas if mesa.zona == 'salon']
        mesas_exteriores = [mesa for mesa in mesas if mesa.zona == 'exteriores']

        self.assertEqual([mesa.etiqueta for mesa in mesas_salon], ['Mesa 1', 'Mesa 2'])
        self.assertEqual(
            [mesa.etiqueta for mesa in mesas_exteriores],
            ['Exteriores 1', 'Exteriores 2', 'Exteriores 3'],
        )
        self.assertEqual(len({mesa.numero for mesa in mesas}), 5)

    def test_aumentar_mesas_de_salon_no_reutiliza_ids_de_mesas_exteriores(self):
        turno = registrar_inicio_turno(self.admin)
        configurar_turno(turno, cantidad_mesas=1, platos=[], cantidad_mesas_exteriores=2)
        configurar_turno(turno, cantidad_mesas=3, cantidad_mesas_exteriores=2)

        mesas = mesas_del_turno(turno)
        mesas_salon = [mesa for mesa in mesas if mesa.zona == 'salon']
        mesas_exteriores = [mesa for mesa in mesas if mesa.zona == 'exteriores']

        self.assertEqual([mesa.etiqueta for mesa in mesas_salon], ['Mesa 1', 'Mesa 2', 'Mesa 3'])
        self.assertEqual([mesa.etiqueta for mesa in mesas_exteriores], ['Exteriores 1', 'Exteriores 2'])
        self.assertEqual(len({mesa.numero for mesa in mesas}), 5)

    def test_administrador_configura_la_cantidad_de_mesas_exteriores(self):
        turno = registrar_inicio_turno(self.admin)
        self.client.force_login(self.admin)

        response = self.client.post(
            reverse('usuarios:configurar_mesas_turno'),
            {'cantidad_mesas': '2', 'cantidad_mesas_exteriores': '2'},
        )

        self.assertRedirects(response, reverse('usuarios:dashboard'))
        turno.refresh_from_db()
        self.assertEqual(turno.cantidad_mesas_exteriores, 2)
        self.assertEqual(
            list(Mesa.objects.filter(zona='exteriores').values_list('numero_zona', flat=True)),
            [1, 2],
        )

    def test_panel_admin_muestra_un_qr_y_estado_por_mesa(self):
        registrar_inicio_turno(self.admin)
        self.client.force_login(self.admin)

        response = self.client.get(reverse('usuarios:dashboard'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Mesas y códigos QR')
        self.assertContains(response, 'Código QR de Mesa 1')
        self.assertContains(response, reverse('usuarios:agregar_mesa_turno'))
        self.assertNotContains(response, 'cantidad-mesas')
        self.assertEqual(len(response.context['mesas_admin']), 10)

    def test_admin_puede_desactivar_y_activar_una_mesa_individualmente(self):
        turno = registrar_inicio_turno(self.admin)
        mesa = Mesa.objects.create(numero=80, activa=True)
        Pedido.objects.create(mesa=mesa, estado='cancelado')
        self.client.force_login(self.admin)

        response = self.client.post(
            reverse('usuarios:cambiar_estado_mesa_turno', args=[mesa.pk]),
            {'activa': '0'},
        )

        self.assertRedirects(response, reverse('usuarios:dashboard'))
        mesa.refresh_from_db()
        self.assertFalse(mesa.activa)
        self.assertNotIn(mesa, mesas_del_turno(turno))

        self.client.post(
            reverse('usuarios:cambiar_estado_mesa_turno', args=[mesa.pk]),
            {'activa': '1'},
        )
        mesa.refresh_from_db()
        self.assertTrue(mesa.activa)
        self.assertIn(mesa, mesas_del_turno(turno))

    def test_mesa_con_cuenta_abierta_no_puede_desactivarse(self):
        registrar_inicio_turno(self.admin)
        mesa = Mesa.objects.create(numero=81, activa=True, abierta=True)
        self.client.force_login(self.admin)

        response = self.client.post(
            reverse('usuarios:cambiar_estado_mesa_turno', args=[mesa.pk]),
            {'activa': '0'},
            follow=True,
        )

        mesa.refresh_from_db()
        self.assertTrue(mesa.activa)
        self.assertContains(response, 'mientras tenga una cuenta abierta')

    def test_admin_puede_agregar_una_mesa_exterior(self):
        turno = registrar_inicio_turno(self.admin)
        self.client.force_login(self.admin)

        response = self.client.post(
            reverse('usuarios:agregar_mesa_turno'),
            {'zona': 'exteriores'},
        )

        self.assertRedirects(response, reverse('usuarios:dashboard'))
        mesa = Mesa.objects.get(zona='exteriores', numero_zona=1)
        self.assertTrue(mesa.activa)
        turno.refresh_from_db()
        self.assertEqual(turno.cantidad_mesas_exteriores, 1)

    def test_turno_expone_metricas_historicas_y_vista_admin(self):
        turno = registrar_inicio_turno(self.mesera)
        registrar_inicio_turno(self.cajera)

        mesa = Mesa.objects.create(numero=10, abierta=True)
        Factura.objects.create(
            pedido=Pedido.objects.create(mesa=mesa, estado='cerrado'),
            mesa_numero=mesa.numero,
            forma_pago='efectivo_cup',
            total_cup=Decimal('200.00'),
            monto_efectivo_cup=Decimal('200.00'),
        )
        Factura.objects.create(
            pedido=Pedido.objects.create(mesa=mesa, estado='cerrado'),
            mesa_numero=mesa.numero,
            forma_pago='transferencia',
            total_cup=Decimal('300.00'),
            monto_transferencia_cup=Decimal('300.00'),
        )

        turno_cerrado = cerrar_turno(self.cajera)

        self.assertEqual(turno_cerrado.total_general_cup, Decimal('500.00'))
        self.assertEqual(turno_cerrado.total_efectivo_cup, Decimal('200.00'))
        self.assertEqual(turno_cerrado.total_transferencia_cup, Decimal('300.00'))

        request = RequestFactory().get('/usuarios/historial/')
        request.user = self.admin
        response = historial_turnos_view(request)

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Historial de turnos')
        self.assertContains(response, '500.00')
        self.assertContains(response, 'General')

    def test_admin_puede_ver_detalle_de_turno_especifico_con_ipv(self):
        turno = registrar_inicio_turno(self.mesera)
        registrar_inicio_turno(self.cajera)

        mesa = Mesa.objects.create(numero=11, abierta=True)
        Factura.objects.create(
            pedido=Pedido.objects.create(mesa=mesa, estado='cerrado'),
            mesa_numero=mesa.numero,
            forma_pago='efectivo_cup',
            total_cup=Decimal('180.00'),
            monto_efectivo_cup=Decimal('180.00'),
        )
        Factura.objects.create(
            pedido=Pedido.objects.create(mesa=mesa, estado='cerrado'),
            mesa_numero=mesa.numero,
            forma_pago='transferencia',
            total_cup=Decimal('220.00'),
            monto_transferencia_cup=Decimal('220.00'),
        )

        turno_cerrado = cerrar_turno(self.cajera)

        request = RequestFactory().get(f'/turnos/{turno_cerrado.pk}/')
        request.user = self.admin
        response = historial_turno_detalle_view(request, turno_cerrado.pk)

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Detalle del turno')
        self.assertContains(response, 'IPV')
        self.assertContains(response, '400.00')
        self.assertContains(
            response,
            reverse('usuarios:ipv_turno_pdf', args=[turno_cerrado.pk]),
        )
        self.assertContains(
            response,
            reverse('usuarios:ipv_turno_excel', args=[turno_cerrado.pk]),
        )

        self.client.force_login(self.cajera)
        response_historial = self.client.get(reverse('usuarios:historial_turnos'))
        self.assertEqual(response_historial.status_code, 302)
        response_detalle = self.client.get(
            reverse('usuarios:historial_turno_detalle', args=[turno_cerrado.pk])
        )
        self.assertEqual(response_detalle.status_code, 302)
        self.assertEqual(
            self.client.get(reverse('usuarios:ipv_turno', args=[turno_cerrado.pk])).status_code,
            403,
        )

    def test_ipv_muestra_existencia_despues_de_restar_salidas(self):
        turno = registrar_inicio_turno(self.admin)
        plato = Plato.objects.create(
            categoria=Categoria.objects.create(nombre='Carta IPV'),
            nombre='Producto IPV',
            precio=Decimal('15.00'),
        )
        ProductoIPVTurno.objects.create(
            turno=turno,
            plato=plato,
            cantidad_inicial=Decimal('10'),
            entradas=Decimal('5'),
            merma_rotura=Decimal('1'),
        )
        mesa = Mesa.objects.create(numero=84)
        Factura.objects.create(
            pedido=Pedido.objects.create(mesa=mesa, estado='cerrado'),
            mesa_numero=mesa.numero,
            forma_pago='efectivo_cup',
            total_cup=Decimal('30.00'),
            items_snapshot=[{
                'plato_id': plato.pk,
                'plato': plato.nombre,
                'cantidad': 2,
                'precio_unit': '15.00',
                'subtotal': '30.00',
            }],
        )

        self.client.force_login(self.admin)
        response = self.client.get(reverse('usuarios:ipv_turno', args=[turno.pk]))

        self.assertEqual(response.status_code, 200)
        producto = response.json()['filas'][0]
        self.assertEqual(producto['nombre'], plato.nombre)
        self.assertEqual(producto['entradas'], '5.000')
        self.assertEqual(producto['cantidad_vendida'], '2')
        self.assertEqual(producto['cantidad_existencia'], '12.000')
        self.assertEqual(producto['merma_rotura'], '1.000')
        self.assertEqual(producto['saldo_final'], '12.000')

    def test_ipv_turno_se_exporta_a_pdf_y_excel(self):
        turno = registrar_inicio_turno(self.admin)
        self.client.force_login(self.admin)
        plato = Plato.objects.create(
            categoria=Categoria.objects.create(nombre='Carta exportación'),
            nombre='=2+3',
            precio=Decimal('125.50'),
        )
        ProductoIPVTurno.objects.create(
            turno=turno,
            plato=plato,
            cantidad_inicial=Decimal('5'),
            entradas=Decimal('3'),
        )

        pdf_response = self.client.get(
            reverse('usuarios:ipv_turno_pdf', args=[turno.pk]),
        )
        self.assertEqual(pdf_response.status_code, 200)
        self.assertEqual(pdf_response['Content-Type'], 'application/pdf')
        self.assertIn(f'ipv_turno_{turno.pk}.pdf', pdf_response['Content-Disposition'])
        self.assertTrue(pdf_response.content.startswith(b'%PDF'))

        excel_response = self.client.get(
            reverse('usuarios:ipv_turno_excel', args=[turno.pk]),
        )
        self.assertEqual(excel_response.status_code, 200)
        self.assertEqual(
            excel_response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertIn(f'ipv_turno_{turno.pk}.xlsx', excel_response['Content-Disposition'])
        libro = load_workbook(BytesIO(excel_response.content), read_only=True)
        hoja = libro['IPV']
        self.assertEqual(hoja['A3'].value, 'Fondo inicial CUP')
        self.assertEqual(hoja['A4'].value, 'Descripción')
        self.assertEqual(hoja['A5'].value, '=2+3')
        self.assertEqual(hoja['A5'].data_type, 's')
        libro.close()

    def test_ipv_exportaciones_funcionan_sin_filas(self):
        turno = registrar_inicio_turno(self.admin)
        self.client.force_login(self.admin)

        pdf_response = self.client.get(reverse('usuarios:ipv_turno_pdf', args=[turno.pk]))
        excel_response = self.client.get(reverse('usuarios:ipv_turno_excel', args=[turno.pk]))

        self.assertEqual(pdf_response.status_code, 200)
        self.assertTrue(pdf_response.content.startswith(b'%PDF'))
        libro = load_workbook(BytesIO(excel_response.content), read_only=True)
        self.assertEqual(libro['IPV']['A3'].value, 'Fondo inicial CUP')
        self.assertEqual(libro['IPV']['A4'].value, 'Descripción')
        libro.close()

    def test_ipv_exportaciones_solo_permiten_roles_financieros(self):
        turno = registrar_inicio_turno(self.admin)
        self.client.force_login(self.mesera)

        for nombre_url in ('ipv_turno_pdf', 'ipv_turno_excel'):
            with self.subTest(exportacion=nombre_url):
                respuesta = self.client.get(reverse(f'usuarios:{nombre_url}', args=[turno.pk]))
                self.assertEqual(respuesta.status_code, 403)

    def test_admin_puede_crear_usuario_con_rol_y_password_autenticable(self):
        self.client.force_login(self.admin)

        response = self.client.post(reverse('usuarios:crear_usuario'), {
            'username': 'cocina_nueva',
            'first_name': 'Ana',
            'last_name': 'García Pérez',
            'password': 'clave-segura-123',
            'role': 'cocina',
        })

        self.assertRedirects(response, reverse('usuarios:gestion_usuarios'))
        usuario = Usuario.objects.get(username='cocina_nueva')
        self.assertEqual(usuario.get_full_name(), 'Ana García Pérez')
        self.assertEqual(usuario.role, 'cocina')
        self.assertTrue(usuario.is_active)
        self.assertTrue(usuario.check_password('clave-segura-123'))
        self.assertTrue(self.client.login(username='cocina_nueva', password='clave-segura-123'))

    def test_login_registra_al_trabajador_con_nombre_completo_en_el_turno(self):
        trabajador = Usuario.objects.create_user(
            username='cocina_turno',
            first_name='Luis',
            last_name='Martínez',
            password='123456',
            role='cocina',
        )
        self.client.force_login(self.admin)
        self.client.post(reverse('usuarios:abrir_turno'))
        self.client.logout()

        response = self.client.post(reverse('usuarios:login'), {
            'username': 'cocina_turno',
            'password': '123456',
        })

        self.assertRedirects(response, reverse('pedidos:cocina'))
        turno = get_turno_abierto()
        self.assertIsNotNone(turno)
        self.assertIn(trabajador, turno.usuarios.all())
        self.assertEqual(trabajador.get_full_name(), 'Luis Martínez')

        self.client.force_login(self.admin)
        response = self.client.get(reverse('usuarios:dashboard'))
        self.assertContains(response, 'Luis Martínez')
        self.assertNotContains(response, 'cocina_turno')
        response_usuarios = self.client.get(reverse('usuarios:gestion_usuarios'))
        self.assertContains(response_usuarios, 'cocina_turno')

    def test_admin_puede_desactivar_usuario_y_se_bloquea_su_login(self):
        self.client.force_login(self.admin)

        response = self.client.post(reverse('usuarios:cambiar_estado_usuario', args=[self.mesera.pk]))

        self.assertRedirects(response, reverse('usuarios:gestion_usuarios'))
        self.mesera.refresh_from_db()
        self.assertFalse(self.mesera.is_active)
        self.assertFalse(self.mesera.activo)
        self.assertFalse(self.client.login(username='mesera1', password='123456'))

    def test_admin_puede_editar_datos_rol_y_password_de_un_usuario(self):
        self.client.force_login(self.admin)

        response = self.client.post(reverse('usuarios:editar_usuario', args=[self.mesera.pk]), {
            'first_name': 'María',
            'last_name': 'López',
            'username': 'mesera_actualizada',
            'email': 'maria@example.com',
            'telefono': '55512345',
            'role': 'cajera',
            'password': 'nueva-clave-456',
        })

        self.assertRedirects(response, reverse('usuarios:gestion_usuarios'))
        self.mesera.refresh_from_db()
        self.assertEqual(self.mesera.get_full_name(), 'María López')
        self.assertEqual(self.mesera.username, 'mesera_actualizada')
        self.assertEqual(self.mesera.role, 'cajera')
        self.assertTrue(self.mesera.check_password('nueva-clave-456'))
        self.assertTrue(self.client.login(username='mesera_actualizada', password='nueva-clave-456'))

    def test_admin_no_puede_quitarse_su_propio_rol(self):
        self.client.force_login(self.admin)

        self.client.post(reverse('usuarios:editar_usuario', args=[self.admin.pk]), {
            'first_name': 'Admin',
            'last_name': 'Principal',
            'username': 'admin',
            'role': 'cocina',
        })

        self.admin.refresh_from_db()
        self.assertEqual(self.admin.role, 'administrador')

    def test_usuario_autenticado_no_carga_el_formulario_de_login(self):
        self.client.force_login(self.mesera)

        response = self.client.get(reverse('usuarios:login'))

        self.assertRedirects(response, reverse('pedidos:mesera'))
        self.assertIn('no-store', response['Cache-Control'])

    def test_login_no_se_guarda_en_cache_y_se_recarga_al_volver_con_atras(self):
        response = self.client.get(reverse('usuarios:landing'))

        self.assertEqual(response.status_code, 200)
        self.assertIn('no-store', response['Cache-Control'])
        self.assertContains(response, "if (event.persisted) window.location.reload();")

        response_login = self.client.get(reverse('usuarios:login'))
        self.assertIn('no-store', response_login['Cache-Control'])

    def test_post_de_login_con_csrf_viejo_redirige_a_la_sesion_activa(self):
        client = Client(enforce_csrf_checks=True)
        client.get(reverse('usuarios:landing'))
        token_viejo = client.cookies['csrftoken'].value

        login_response = client.post(reverse('usuarios:login'), {
            'username': self.mesera.username,
            'password': '123456',
            'csrfmiddlewaretoken': token_viejo,
        })

        self.assertEqual(login_response.status_code, 302)
        token_actual = client.cookies['csrftoken'].value
        self.assertNotEqual(token_actual, token_viejo)
        respuesta_token_viejo = client.post(reverse('usuarios:login'), {
            'username': self.admin.username,
            'password': '123456',
            'csrfmiddlewaretoken': token_viejo,
        })

        self.assertEqual(respuesta_token_viejo.status_code, 403)
        self.assertContains(respuesta_token_viejo, 'Formulario vencido', status_code=403)
        self.assertContains(respuesta_token_viejo, 'Tu sesión actual sigue activa', status_code=403)
        self.assertContains(respuesta_token_viejo, reverse('pedidos:mesera'), status_code=403)
        self.assertIn('no-store', respuesta_token_viejo['Cache-Control'])
        self.assertTrue(client.session.get('_auth_user_id'))
