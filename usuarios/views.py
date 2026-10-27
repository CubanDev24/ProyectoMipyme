from collections import defaultdict
from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO
from xml.sax.saxutils import escape

from django.contrib import messages
from django.contrib.auth import logout, login, authenticate, get_user, update_session_auth_hash
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Max, Q, Sum
from django.db.models.functions import Coalesce
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_POST, require_http_methods
from django.views.csrf import csrf_failure as django_csrf_failure
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from carta.models import Plato
from carta.qr import generar_qr_mesa
from inventario.models import Insumo, MovimientoInventario
from pedidos.models import Factura, Mesa, Pedido
from .models import Notificacion, ProductoIPVTurno, Turno, Usuario, cerrar_turno, configurar_turno, crear_mesas_del_turno, get_turno_abierto


ROLE_HOME = {
    'administrador': 'inventario:administrador',
    'cocina': 'pedidos:cocina',
    'mesera': 'pedidos:mesera',
    'cajera': 'pedidos:caja',
}
ROLES_CON_ACCESO_FINANCIERO = {'administrador', 'cajera'}


@never_cache
def csrf_failure(request, reason=''):
    if request.path != reverse('usuarios:login') or request.method != 'POST':
        return django_csrf_failure(request, reason=reason)

    user = get_user(request)
    autenticado = user.is_authenticated
    if autenticado:
        destino = reverse(ROLE_HOME.get(user.role, 'usuarios:dashboard'))
    else:
        destino = reverse('usuarios:landing')
    return render(request, 'usuarios/csrf_failure.html', {
        'autenticado': autenticado,
        'destino': destino,
    }, status=403)


@never_cache
def landing(request):
    if request.user.is_authenticated:
        url = ROLE_HOME.get(request.user.role, 'carta:index')
        return redirect(url)
    return render(request, 'usuarios/login.html')


@never_cache
@require_http_methods(['GET', 'POST'])
def login_view(request):
    if request.user.is_authenticated:
        return redirect(ROLE_HOME.get(request.user.role, 'usuarios:dashboard'))

    if request.method == 'GET':
        return redirect('usuarios:landing')

    username = request.POST.get('username', '').strip()
    password = request.POST.get('password', '')
    user = authenticate(request, username=username, password=password)
    if user is None:
        messages.error(request, 'Credenciales inválidas.')
        return redirect('usuarios:landing')

    login(request, user)
    turno = get_turno_abierto()
    if turno:
        turno.usuarios.add(user)
        user.ultimo_login_turno = timezone.now()
        user.save(update_fields=['ultimo_login_turno'])
    return redirect(ROLE_HOME.get(user.role, 'usuarios:landing'))


@login_required
def dashboard(request):
    hoy = timezone.localdate()
    puede_ver_datos_financieros = request.user.role in ROLES_CON_ACCESO_FINANCIERO
    turno = get_turno_abierto()
    turno_anterior = None
    turno_abierto_pendiente = None
    if request.user.role == 'administrador':
        turno = turno or Turno.objects.filter(fecha=hoy).prefetch_related('usuarios').order_by('-apertura').first()
        turno_anterior = Turno.objects.filter(
            estado='cerrado', fecha__lt=hoy,
        ).prefetch_related('usuarios').order_by('-fecha', '-apertura').first()
        turno_abierto_pendiente = Turno.objects.filter(
            estado='abierto', fecha__lt=hoy, continuidad_confirmada=False,
        ).order_by('-fecha', '-apertura').first()
    turno_requiere_confirmacion = bool(
        request.user.role == 'administrador'
        and turno
        and turno_abierto_pendiente
        and turno.pk == turno_abierto_pendiente.pk
    )
    mesas_admin = []
    mesas_activas_salon = 0
    mesas_activas_exteriores = 0
    if turno:
        if request.user.role == 'administrador' and turno.estado == 'abierto':
            crear_mesas_del_turno(turno)
        mesas_activas = Mesa.objects.filter(activa=True).values_list('zona', flat=True)
        for zona in mesas_activas:
            if zona == 'exteriores':
                mesas_activas_exteriores += 1
            else:
                mesas_activas_salon += 1
        if request.user.role == 'administrador' and turno.estado == 'abierto':
            mesas_admin = list(Mesa.objects.order_by('zona', 'numero_zona', 'numero'))
        for mesa in mesas_admin:
            mesa.qr_url, mesa.qr_data = generar_qr_mesa(request, mesa)
    notificaciones = Notificacion.objects.filter(destinatario=request.user).order_by('-creada_en')[:10]
    ipv_configurado = False
    opciones_ipv = []
    if turno and turno.estado == 'abierto' and request.user.role in ROLES_CON_ACCESO_FINANCIERO:
        ipv_configurado = ProductoIPVTurno.objects.filter(turno=turno).exists()
        if not ipv_configurado:
            platos = list(
                Plato.objects.filter(disponible=True)
                .select_related('categoria')
                .order_by('categoria__orden', 'nombre')
            )
            cantidades_iniciales = _cantidades_finales_productos_ipv(platos)
            registros = [
                ProductoIPVTurno(
                    turno=turno,
                    plato=plato,
                    cantidad_inicial=cantidades_iniciales.get(plato.pk, Decimal('0')),
                    entradas=Decimal('0'),
                )
                for plato in platos
            ]
            ProductoIPVTurno.objects.bulk_create(registros)
            ipv_configurado = True
        opciones_ipv = [
            {
                'plato': registro.plato,
                'cantidad_inicial': registro.cantidad_inicial,
            }
            for registro in ProductoIPVTurno.objects.filter(turno=turno)
            .select_related('plato__categoria')
            .order_by('plato__nombre')
        ]
    facturas_turno = Factura.objects.none()
    if turno and puede_ver_datos_financieros:
        fin_turno = turno.cierre or timezone.now()
        facturas_turno = Factura.objects.filter(
            creado_en__gte=turno.apertura,
            creado_en__lte=fin_turno,
        ).select_related('pedido__mesa').order_by('-creado_en')
    return render(request, 'usuarios/dashboard.html', {
        'turno': turno,
        'hoy': hoy,
        'turno_anterior': turno_anterior,
        'resumen_turno_anterior': turno_anterior.resumen_financiero if turno_anterior and puede_ver_datos_financieros else None,
        'resumen_turno': turno.resumen_financiero if turno and puede_ver_datos_financieros else None,
        'puede_ver_datos_financieros': puede_ver_datos_financieros,
        'turno_abierto_pendiente': turno_abierto_pendiente,
        'turno_requiere_confirmacion': turno_requiere_confirmacion,
        'mesas_admin': mesas_admin,
        'mesas_activas_salon': mesas_activas_salon,
        'mesas_activas_exteriores': mesas_activas_exteriores,
        'facturas_turno': facturas_turno,
        'notificaciones': notificaciones,
        'ipv_configurado': ipv_configurado,
        'opciones_ipv': opciones_ipv,
    })


@login_required
def gestion_usuarios_view(request):
    if request.user.role != 'administrador':
        messages.error(request, 'Solo el administrador puede gestionar usuarios.')
        return redirect('usuarios:dashboard')
    return render(request, 'usuarios/gestion_usuarios.html', {
        'usuarios': Usuario.objects.order_by('role', 'username'),
    })


@login_required
@require_http_methods(['POST'])
def abrir_turno_view(request):
    if request.user.role != 'administrador':
        messages.error(request, 'Solo el administrador puede abrir el turno.')
        return redirect('usuarios:dashboard')

    hoy = timezone.localdate()
    if Turno.objects.filter(fecha=hoy).exists():
        messages.error(request, 'Ya existe un turno registrado para hoy.')
        return redirect('usuarios:dashboard')
    if Turno.objects.filter(estado='abierto').exists():
        messages.error(request, 'Hay un turno anterior abierto. Debe cerrarse antes de abrir uno nuevo.')
        return redirect('usuarios:dashboard')

    turno = Turno.objects.create(fecha=hoy, estado='abierto')
    turno.usuarios.add(request.user)
    request.user.ultimo_login_turno = timezone.now()
    request.user.save(update_fields=['ultimo_login_turno'])
    crear_mesas_del_turno(turno)
    messages.success(request, f'Turno del {hoy:%d/%m/%Y} abierto correctamente.')
    return redirect('usuarios:dashboard')


@login_required
@require_POST
def continuar_turno_view(request):
    if request.user.role != 'administrador':
        messages.error(request, 'Solo el administrador puede confirmar la continuidad del turno.')
        return redirect('usuarios:dashboard')

    turno = get_object_or_404(
        Turno,
        pk=request.POST.get('turno_id'),
        estado='abierto',
        fecha__lt=timezone.localdate(),
    )
    turno.continuidad_confirmada = True
    turno.save(update_fields=['continuidad_confirmada'])
    messages.success(request, 'El turno continuará abierto hasta que se realice su cierre.')
    return redirect('usuarios:dashboard')


def _decimal_o_cero(valor):
    try:
        return Decimal(str(valor).replace(',', '.').strip())
    except (ArithmeticError, TypeError, ValueError, AttributeError):
        return Decimal('0')


def _decimal_ipv_no_negativo(valor, nombre_campo, decimal_places=3):
    try:
        cantidad = Decimal(str(valor).replace(',', '.').strip())
    except (ArithmeticError, TypeError, ValueError, AttributeError):
        raise ValueError(f'Introduce una cantidad válida para {nombre_campo}.')
    if not cantidad.is_finite() or cantidad < 0:
        raise ValueError(f'{nombre_campo.capitalize()} no puede ser negativo.')
    if cantidad.as_tuple().exponent < -decimal_places:
        raise ValueError(f'{nombre_campo.capitalize()} admite como máximo {decimal_places} decimales.')
    if cantidad and cantidad.adjusted() >= 12 - decimal_places:
        raise ValueError(f'{nombre_campo.capitalize()} supera el límite permitido.')
    return cantidad


def _ventas_facturadas_por_plato(turno, cierre, platos):
    platos_por_id = {plato.pk: plato for plato in platos}
    ids_por_nombre = {}
    for plato in platos:
        ids_por_nombre.setdefault(plato.nombre.strip().casefold(), []).append(plato.pk)

    ventas = {
        plato.pk: {
            'cantidad': Decimal('0'),
            'importe': Decimal('0'),
            'precio_total': Decimal('0'),
        }
        for plato in platos
    }
    snapshots = Factura.objects.filter(
        creado_en__gte=turno.apertura,
        creado_en__lte=cierre,
    ).values_list('items_snapshot', flat=True)
    for snapshot in snapshots:
        for item in snapshot or []:
            plato_id = item.get('plato_id')
            try:
                plato_id = int(plato_id) if plato_id is not None else None
            except (TypeError, ValueError):
                plato_id = None
            if plato_id not in platos_por_id:
                nombre = str(item.get('plato', '')).strip().casefold()
                candidatos = ids_por_nombre.get(nombre, [])
                plato_id = candidatos[0] if len(candidatos) == 1 else None
            if plato_id is None:
                continue

            cantidad = _decimal_o_cero(item.get('cantidad'))
            datos = ventas[plato_id]
            datos['cantidad'] += cantidad
            datos['importe'] += _decimal_o_cero(item.get('subtotal'))
            datos['precio_total'] += _decimal_o_cero(item.get('precio_unit')) * cantidad
    return ventas


def _cantidades_finales_productos_ipv(platos):
    if not platos:
        return {}
    registros = list(ProductoIPVTurno.objects.filter(
        plato_id__in=[plato.pk for plato in platos],
        turno__estado='cerrado',
        turno__cierre__isnull=False,
    ).select_related('turno', 'plato').order_by(
        'plato_id', '-turno__cierre', '-turno__apertura',
    ))
    ultimos_por_plato = {}
    for registro in registros:
        ultimos_por_plato.setdefault(registro.plato_id, registro)

    registros_por_turno = defaultdict(list)
    for registro in ultimos_por_plato.values():
        registros_por_turno[registro.turno_id].append(registro)

    cantidades = {}
    for registros_turno in registros_por_turno.values():
        turno_anterior = registros_turno[0].turno
        ventas = _ventas_facturadas_por_plato(
            turno_anterior,
            turno_anterior.cierre,
            [registro.plato for registro in registros_turno],
        )
        for registro in registros_turno:
            cantidades[registro.plato_id] = (
                registro.cantidad_inicial
                + registro.entradas
                - registro.merma_rotura
                - ventas[registro.plato_id]['cantidad']
            )
    return cantidades


@login_required
@require_POST
def configurar_ipv_turno_view(request):
    if request.user.role not in ROLES_CON_ACCESO_FINANCIERO:
        return JsonResponse({'error': 'Sin permiso'}, status=403)

    turno = get_turno_abierto()
    if turno is None:
        messages.error(request, 'No hay un turno abierto para configurar el IPV.')
        return redirect('usuarios:dashboard')

    plato_ids = list(dict.fromkeys(request.POST.getlist('platos')))
    productos_ipv_existentes = ProductoIPVTurno.objects.filter(turno=turno).exists()
    platos = []
    if not productos_ipv_existentes:
        platos = list(Plato.objects.filter(pk__in=plato_ids, disponible=True).order_by('pk'))
        if not platos:
            messages.error(request, 'Selecciona al menos un producto de la carta.')
            return redirect('usuarios:dashboard')
        if len(platos) != len(plato_ids):
            messages.error(request, 'Uno o más productos seleccionados ya no están disponibles.')
            return redirect('usuarios:dashboard')

    try:
        fondo_inicial = _decimal_ipv_no_negativo(
            request.POST.get('fondo_inicial_cup'),
            'el fondo inicial',
            decimal_places=2,
        )
        cantidades_iniciales = _cantidades_finales_productos_ipv(platos)
        registros = [
            ProductoIPVTurno(
                turno=turno,
                plato=plato,
                cantidad_inicial=cantidades_iniciales.get(plato.pk, Decimal('0')),
                entradas=Decimal('0'),
            )
            for plato in platos
        ]
    except ValueError as exc:
        messages.error(request, str(exc))
        return redirect('usuarios:dashboard')

    with transaction.atomic():
        turno = Turno.objects.select_for_update().filter(pk=turno.pk, estado='abierto').first()
        if turno is None:
            messages.error(request, 'El turno se cerró antes de configurar su IPV.')
            return redirect('usuarios:dashboard')
        if not ProductoIPVTurno.objects.filter(turno=turno).exists() and registros:
            ProductoIPVTurno.objects.bulk_create(registros)
        turno.fondo_inicial_cup = fondo_inicial
        turno.save(update_fields=['fondo_inicial_cup'])

    messages.success(request, 'El IPV del turno quedó configurado.')
    return redirect('usuarios:dashboard')


@login_required
@require_POST
def registrar_merma_ipv_view(request, producto_turno_id):
    if request.user.role not in ROLES_CON_ACCESO_FINANCIERO:
        return JsonResponse({'error': 'Sin permiso'}, status=403)
    try:
        cantidad = _decimal_ipv_no_negativo(request.POST.get('cantidad'), 'la rotura')
        if cantidad <= 0:
            raise ValueError('La cantidad de rotura debe ser mayor que cero.')
        with transaction.atomic():
            registro = ProductoIPVTurno.objects.select_for_update().select_related('turno').get(
                pk=producto_turno_id,
                turno__estado='abierto',
            )
            cierre = timezone.now()
            venta = _ventas_facturadas_por_plato(
                registro.turno,
                cierre,
                [registro.plato],
            )[registro.plato_id]
            disponible = (
                registro.cantidad_inicial
                + registro.entradas
                - registro.merma_rotura
                - venta['cantidad']
            )
            if cantidad > disponible:
                raise ValueError('La rotura no puede superar las unidades disponibles.')
            registro.merma_rotura += cantidad
            registro.save(update_fields=['merma_rotura'])
    except ProductoIPVTurno.DoesNotExist:
        return JsonResponse({'error': 'El producto ya no pertenece a un turno abierto.'}, status=404)
    except ValueError as exc:
        return JsonResponse({'error': str(exc)}, status=400)
    return JsonResponse({'ok': True, 'merma_rotura': str(registro.merma_rotura)})


@login_required
def inventario_ventas_turno(request):
    if request.user.role not in ROLES_CON_ACCESO_FINANCIERO:
        return JsonResponse({'error': 'Sin permiso'}, status=403)

    turno = get_turno_abierto()
    if turno is None:
        return JsonResponse({
            'turno_abierto': False,
            'productos': [],
            'actualizado_en': timezone.localtime().strftime('%H:%M:%S'),
        })
    if not ProductoIPVTurno.objects.filter(turno=turno).exists():
        platos = list(
            Plato.objects.filter(disponible=True)
            .select_related('categoria')
            .order_by('categoria__orden', 'nombre')
        )
        cantidades_iniciales = _cantidades_finales_productos_ipv(platos)
        ProductoIPVTurno.objects.bulk_create([
            ProductoIPVTurno(
                turno=turno,
                plato=plato,
                cantidad_inicial=cantidades_iniciales.get(plato.pk, Decimal('0')),
                entradas=Decimal('0'),
            )
            for plato in platos
        ])
    datos = _datos_ipv_turno(turno)
    return JsonResponse({
        **datos,
        'turno_abierto': True,
        'productos': datos['filas'],
    })


@login_required
def logout_view(request):
    logout(request)
    return redirect('usuarios:landing')


@login_required
@require_http_methods(['POST'])
def crear_usuario(request):
    if request.user.role != 'administrador':
        messages.error(request, 'Solo el administrador puede crear usuarios.')
        return redirect('usuarios:dashboard')

    username = request.POST.get('username', '').strip()
    first_name = request.POST.get('first_name', '').strip()
    last_name = request.POST.get('last_name', '').strip()
    password = request.POST.get('password', '')
    role = request.POST.get('role')
    if not username or not first_name or not last_name or not password or role not in dict(Usuario.ROLE_CHOICES):
        messages.error(request, 'Completa los datos del usuario.')
        return redirect('usuarios:gestion_usuarios')

    if Usuario.objects.filter(username__iexact=username).exists():
        messages.error(request, f'El usuario {username} ya existe.')
        return redirect('usuarios:gestion_usuarios')

    Usuario.objects.create_user(
        username=username,
        first_name=first_name,
        last_name=last_name,
        password=password,
        role=role,
        is_active=True,
        activo=True,
    )
    messages.success(request, f'Usuario {first_name} {last_name} creado con rol {dict(Usuario.ROLE_CHOICES)[role]}.')
    return redirect('usuarios:gestion_usuarios')


@login_required
@require_http_methods(['POST'])
def cambiar_estado_usuario(request, user_id):
    if request.user.role != 'administrador':
        messages.error(request, 'Solo el administrador puede gestionar usuarios.')
        return redirect('usuarios:gestion_usuarios')

    usuario = get_object_or_404(Usuario, pk=user_id)
    if usuario == request.user:
        messages.error(request, 'No puedes desactivar tu propio usuario.')
        return redirect('usuarios:gestion_usuarios')

    usuario.is_active = not usuario.is_active
    usuario.activo = usuario.is_active
    usuario.save(update_fields=['is_active', 'activo'])
    estado = 'activado' if usuario.is_active else 'desactivado'
    messages.success(request, f'Usuario {usuario.username} {estado}.')
    return redirect('usuarios:gestion_usuarios')


@login_required
@require_http_methods(['POST'])
def editar_usuario(request, user_id):
    if request.user.role != 'administrador':
        messages.error(request, 'Solo el administrador puede gestionar usuarios.')
        return redirect('usuarios:gestion_usuarios')

    usuario = get_object_or_404(Usuario, pk=user_id)
    username = request.POST.get('username', '').strip()
    first_name = request.POST.get('first_name', '').strip()
    last_name = request.POST.get('last_name', '').strip()
    role = request.POST.get('role')
    password = request.POST.get('password', '')
    if not username or not first_name or not last_name or role not in dict(Usuario.ROLE_CHOICES):
        messages.error(request, 'Completa nombre, apellidos, usuario y rol.')
        return redirect('usuarios:gestion_usuarios')

    if Usuario.objects.filter(username__iexact=username).exclude(pk=usuario.pk).exists():
        messages.error(request, f'El usuario {username} ya existe.')
        return redirect('usuarios:gestion_usuarios')
    if usuario == request.user and role != 'administrador':
        messages.error(request, 'No puedes quitarte el rol de administrador a ti mismo.')
        return redirect('usuarios:gestion_usuarios')

    usuario.username = username
    usuario.first_name = first_name
    usuario.last_name = last_name
    usuario.email = request.POST.get('email', '').strip()
    usuario.telefono = request.POST.get('telefono', '').strip()
    usuario.role = role
    if password:
        usuario.set_password(password)
    usuario.save()
    if usuario == request.user and password:
        update_session_auth_hash(request, usuario)
    messages.success(request, f'Usuario {usuario.get_full_name()} actualizado correctamente.')
    return redirect('usuarios:gestion_usuarios')


@login_required
@require_http_methods(['POST'])
def configurar_mesas_turno_view(request):
    if request.user.role != 'administrador':
        messages.error(request, 'Solo el administrador puede configurar las mesas del turno.')
        return redirect('usuarios:dashboard')

    turno = get_turno_abierto()
    if turno is None:
        messages.error(request, 'Primero debes abrir el turno del día.')
        return redirect('usuarios:dashboard')

    cantidad_mesas = request.POST.get('cantidad_mesas', '').strip()
    cantidad_mesas_exteriores = request.POST.get('cantidad_mesas_exteriores', '0').strip()
    if not cantidad_mesas:
        messages.error(request, 'Debes indicar la cantidad de mesas.')
        return redirect('usuarios:dashboard')

    try:
        cantidad = int(cantidad_mesas)
        cantidad_exteriores = int(cantidad_mesas_exteriores)
    except ValueError:
        messages.error(request, 'Las cantidades de mesas deben ser números enteros.')
        return redirect('usuarios:dashboard')

    if cantidad_exteriores < 0:
        messages.error(request, 'La cantidad de mesas exteriores no puede ser negativa.')
        return redirect('usuarios:dashboard')

    turno.cantidad_mesas = max(cantidad, 1)
    turno.cantidad_mesas_exteriores = cantidad_exteriores
    turno.save(update_fields=['cantidad_mesas', 'cantidad_mesas_exteriores'])
    from usuarios.models import crear_mesas_del_turno
    crear_mesas_del_turno(turno)
    messages.success(
        request,
        f'Mesas actualizadas: {turno.cantidad_mesas} de salón y '
        f'{turno.cantidad_mesas_exteriores} exteriores.',
    )
    return redirect('usuarios:dashboard')


@login_required
@require_http_methods(['POST'])
def agregar_mesa_turno_view(request):
    if request.user.role != 'administrador':
        messages.error(request, 'Solo el administrador puede gestionar las mesas del turno.')
        return redirect('usuarios:dashboard')

    turno = get_turno_abierto()
    if turno is None:
        messages.error(request, 'Primero debes abrir el turno del día.')
        return redirect('usuarios:dashboard')

    zona = request.POST.get('zona', '')
    if zona not in ('salon', 'exteriores'):
        messages.error(request, 'Selecciona una zona válida para la mesa.')
        return redirect('usuarios:dashboard')

    siguiente_numero_zona = (
        Mesa.objects.filter(zona=zona).aggregate(
            ultimo=Max(Coalesce('numero_zona', 'numero')),
        )['ultimo'] or 0
    ) + 1
    siguiente_numero = (Mesa.objects.aggregate(ultimo=Max('numero'))['ultimo'] or 0) + 1
    mesa = Mesa.objects.create(
        numero=siguiente_numero,
        zona=zona,
        numero_zona=siguiente_numero_zona,
        activa=True,
        abierta=False,
    )
    if zona == 'exteriores':
        turno.cantidad_mesas_exteriores = max(turno.cantidad_mesas_exteriores, siguiente_numero_zona)
        turno.save(update_fields=['cantidad_mesas_exteriores'])
    else:
        turno.cantidad_mesas = max(turno.cantidad_mesas, siguiente_numero_zona)
        turno.save(update_fields=['cantidad_mesas'])
    messages.success(request, f'{mesa.etiqueta} agregada y activada.')
    return redirect('usuarios:dashboard')


@login_required
@require_http_methods(['POST'])
def cambiar_estado_mesa_turno_view(request, mesa_id):
    if request.user.role != 'administrador':
        messages.error(request, 'Solo el administrador puede gestionar las mesas del turno.')
        return redirect('usuarios:dashboard')

    if get_turno_abierto() is None:
        messages.error(request, 'Primero debes abrir el turno del día.')
        return redirect('usuarios:dashboard')

    mesa = get_object_or_404(Mesa, pk=mesa_id)
    estado = request.POST.get('activa')
    if estado not in ('0', '1'):
        messages.error(request, 'El estado indicado para la mesa no es válido.')
        return redirect('usuarios:dashboard')

    activar = estado == '1'
    if not activar and (mesa.abierta or Pedido.objects.filter(mesa=mesa).exclude(estado='cerrado').exists()):
        messages.error(request, f'No se puede desactivar {mesa.etiqueta} mientras tenga una cuenta abierta.')
        return redirect('usuarios:dashboard')

    mesa.activa = activar
    mesa.save(update_fields=['activa'])
    estado_texto = 'activada' if activar else 'desactivada'
    messages.success(request, f'{mesa.etiqueta} {estado_texto}.')
    return redirect('usuarios:dashboard')


@login_required
@require_http_methods(['POST'])
def cerrar_turno_view(request):
    if request.user.role not in ('cajera', 'administrador'):
        messages.error(request, 'Solo caja o el administrador pueden cerrar el turno.')
        return redirect('usuarios:dashboard')

    try:
        turno_pendiente = None
        turno_id = request.POST.get('turno_id')
        if turno_id and request.user.role != 'administrador':
            messages.error(request, 'Solo el administrador puede cerrar un turno anterior.')
            return redirect('usuarios:dashboard')
        if turno_id:
            turno_pendiente = get_object_or_404(
                Turno,
                pk=turno_id,
                estado='abierto',
                fecha__lt=timezone.localdate(),
            )
        turno = cerrar_turno(
            request.user,
            observaciones=request.POST.get('observaciones', ''),
            turno=turno_pendiente,
        )
        messages.success(request, f'Turno cerrado correctamente. Resumen: {turno.resumen}')
        if request.user.role == 'administrador':
            return redirect('usuarios:historial_turno_detalle', turno_id=turno.pk)
    except PermissionError as exc:
        messages.error(request, str(exc))
    except ValueError as exc:
        messages.error(request, str(exc))
    return redirect('usuarios:dashboard')


@login_required
def historial_turnos_view(request):
    if request.user.role != 'administrador':
        messages.error(request, 'Solo el administrador puede ver el historial de turnos.')
        return redirect('usuarios:dashboard')

    turnos = Turno.objects.filter(estado='cerrado').prefetch_related('usuarios').order_by('-cierre', '-fecha')
    texto = request.GET.get('texto', '').strip()
    mes = request.GET.get('mes', '').strip()
    anio = request.GET.get('anio', '').strip()
    fecha_desde = request.GET.get('fecha_desde', '').strip()
    fecha_hasta = request.GET.get('fecha_hasta', '').strip()
    error_filtros = ''

    fecha_desde_valor = parse_date(fecha_desde) if fecha_desde else None
    fecha_hasta_valor = parse_date(fecha_hasta) if fecha_hasta else None
    if (fecha_desde and fecha_desde_valor is None) or (fecha_hasta and fecha_hasta_valor is None):
        error_filtros = 'Introduce un rango de fechas válido.'
    elif fecha_desde_valor and fecha_hasta_valor and fecha_desde_valor > fecha_hasta_valor:
        error_filtros = 'La fecha inicial debe ser anterior o igual a la fecha final.'

    if error_filtros:
        turnos = turnos.none()
    else:
        if fecha_desde_valor:
            turnos = turnos.filter(fecha__gte=fecha_desde_valor)
        if fecha_hasta_valor:
            turnos = turnos.filter(fecha__lte=fecha_hasta_valor)

        if mes:
            if mes.isdecimal() and 1 <= int(mes) <= 12:
                turnos = turnos.filter(fecha__month=int(mes))
            else:
                error_filtros = 'El mes seleccionado no es válido.'
                turnos = turnos.none()
        if anio:
            if anio.isdecimal() and len(anio) == 4:
                turnos = turnos.filter(fecha__year=int(anio))
            else:
                error_filtros = 'El año seleccionado no es válido.'
                turnos = turnos.none()

        if texto:
            busqueda = (
                Q(usuarios__username__icontains=texto)
                | Q(usuarios__first_name__icontains=texto)
                | Q(usuarios__last_name__icontains=texto)
                | Q(resumen__icontains=texto)
                | Q(observaciones__icontains=texto)
            )
            if texto.isdecimal():
                busqueda |= Q(pk=int(texto))
            turnos = turnos.filter(busqueda).distinct()

    historial = []
    total_general = Decimal('0')
    total_efectivo = Decimal('0')
    total_transferencia = Decimal('0')
    total_usd = Decimal('0')
    total_facturas = 0
    total_horas = Decimal('0')

    for turno in turnos:
        data = turno.resumen_financiero
        historial.append({'turno': turno, **data})
        total_general += data['total_general_cup']
        total_efectivo += data['total_efectivo_cup']
        total_transferencia += data['total_transferencia_cup']
        total_usd += data['total_usd']
        total_facturas += data['cantidad_facturas']
        total_horas += data['horas_abiertas']

    promedio_ipv = (total_general / Decimal(total_facturas)) if total_facturas else Decimal('0')
    rendimiento_promedio = (total_general / total_horas) if total_horas > 0 else Decimal('0')
    turnos_paginados = Paginator(historial, 10).get_page(request.GET.get('page'))
    pagination_query = request.GET.copy()
    pagination_query.pop('page', None)
    anios_disponibles = [
        fecha.year for fecha in Turno.objects.filter(estado='cerrado').dates('fecha', 'year', order='DESC')
    ]

    context = {
        'historial': turnos_paginados,
        'historial_pagination_query': pagination_query.urlencode(),
        'texto_filtro': texto,
        'mes_filtro': mes,
        'anio_filtro': anio,
        'fecha_desde_filtro': fecha_desde,
        'fecha_hasta_filtro': fecha_hasta,
        'anios_disponibles': anios_disponibles,
        'error_filtros': error_filtros,
        'filtros_aplicados': any((texto, mes, anio, fecha_desde, fecha_hasta)),
        'resumen_general': {
            'total_general_cup': total_general,
            'total_efectivo_cup': total_efectivo,
            'total_transferencia_cup': total_transferencia,
            'total_usd': total_usd,
            'total_facturas': total_facturas,
            'promedio_ipv': promedio_ipv,
            'rendimiento_promedio': rendimiento_promedio,
            'turnos_cerrados': len(historial),
        },
    }
    return render(request, 'usuarios/historial_turnos.html', context)


@login_required
def reportes_view(request):
    if request.user.role != 'administrador':
        messages.error(request, 'Solo el administrador puede consultar los reportes.')
        return redirect('usuarios:dashboard')

    hoy = timezone.localdate()
    periodo = request.GET.get('periodo', 'semana')
    fecha_desde_raw = request.GET.get('fecha_desde', '').strip()
    fecha_hasta_raw = request.GET.get('fecha_hasta', '').strip()
    error_fechas = ''

    if periodo == 'mes':
        fecha_desde = hoy.replace(day=1)
        fecha_hasta = hoy
        periodo_label = 'Mes actual'
    elif periodo == 'mes_anterior':
        inicio_mes_actual = hoy.replace(day=1)
        fecha_hasta = inicio_mes_actual - timedelta(days=1)
        fecha_desde = fecha_hasta.replace(day=1)
        periodo_label = 'Mes anterior'
    elif periodo == 'ano':
        fecha_desde = hoy.replace(month=1, day=1)
        fecha_hasta = hoy
        periodo_label = 'Año actual'
    elif periodo == 'rango':
        try:
            fecha_desde = parse_date(fecha_desde_raw) if fecha_desde_raw else None
            fecha_hasta = parse_date(fecha_hasta_raw) if fecha_hasta_raw else None
        except ValueError:
            fecha_desde = None
            fecha_hasta = None
        periodo_label = 'Rango personalizado'
        if not fecha_desde or not fecha_hasta:
            error_fechas = 'Selecciona las fechas inicial y final para consultar el periodo.'
        elif fecha_desde > fecha_hasta:
            error_fechas = 'La fecha inicial debe ser anterior o igual a la fecha final.'
    else:
        periodo = 'semana'
        fecha_desde = hoy - timedelta(days=hoy.weekday())
        fecha_hasta = hoy
        periodo_label = 'Semana actual'

    facturas = Factura.objects.none()
    if not error_fechas:
        desde = timezone.make_aware(datetime.combine(fecha_desde, datetime.min.time()))
        hasta = timezone.make_aware(datetime.combine(fecha_hasta, datetime.max.time()))
        facturas = Factura.objects.filter(creado_en__gte=desde, creado_en__lte=hasta).order_by('creado_en')

    categorias_por_producto = {
        nombre.strip().casefold(): categoria or 'Sin categoría'
        for nombre, categoria in Plato.objects.values_list('nombre', 'categoria__nombre')
    }
    cantidades_productos = defaultdict(int)
    cantidades_categorias = defaultdict(int)
    recaudacion_por_forma = {
        forma: {
            'nombre': nombre,
            'facturas': 0,
            'total_cup': Decimal('0'),
            'total_usd': Decimal('0'),
        }
        for forma, nombre in Factura.FORMA_PAGO_CHOICES
    }
    recaudacion_por_periodo = defaultdict(lambda: Decimal('0'))
    total_cup = Decimal('0')
    total_usd = Decimal('0')
    total_efectivo_cup = Decimal('0')
    total_transferencia_cup = Decimal('0')
    total_unidades = 0
    items_invalidos = 0
    cantidad_facturas = 0

    for factura in facturas.values(
        'forma_pago',
        'total_cup',
        'monto_usd',
        'monto_efectivo_cup',
        'monto_transferencia_cup',
        'items_snapshot',
        'creado_en',
    ).iterator(chunk_size=500):
        cantidad_facturas += 1
        total_cup += factura['total_cup']
        total_usd += factura['monto_usd']
        total_efectivo_cup += factura['monto_efectivo_cup']
        total_transferencia_cup += factura['monto_transferencia_cup']

        forma_pago = factura['forma_pago']
        if forma_pago in recaudacion_por_forma:
            resumen_pago = recaudacion_por_forma[forma_pago]
            resumen_pago['facturas'] += 1
            resumen_pago['total_cup'] += factura['total_cup']
            resumen_pago['total_usd'] += factura['monto_usd']

        fecha_factura = timezone.localtime(factura['creado_en']).date()
        clave_periodo = fecha_factura.replace(day=1) if (fecha_hasta - fecha_desde).days > 30 else fecha_factura
        recaudacion_por_periodo[clave_periodo] += factura['total_cup']

        for item in factura['items_snapshot'] or []:
            if not isinstance(item, dict):
                items_invalidos += 1
                continue
            try:
                cantidad = int(item.get('cantidad', 1))
            except (TypeError, ValueError):
                items_invalidos += 1
                continue
            if cantidad <= 0:
                items_invalidos += 1
                continue
            nombre_producto = str(item.get('plato') or 'Producto sin nombre').strip()
            nombre_categoria = str(
                item.get('categoria')
                or categorias_por_producto.get(nombre_producto.casefold(), 'Sin categoría')
            ).strip()
            cantidades_productos[nombre_producto] += cantidad
            cantidades_categorias[nombre_categoria] += cantidad
            total_unidades += cantidad

    colores_grafico = ['#34d399', '#38bdf8', '#a78bfa', '#fbbf24', '#fb7185', '#2dd4bf', '#818cf8', '#f97316']
    max_producto = max(cantidades_productos.values(), default=0)
    max_categoria = max(cantidades_categorias.values(), default=0)
    productos_mas_vendidos = [
        {
            'nombre': nombre,
            'cantidad': cantidad,
            'porcentaje': max(2, round(cantidad * 100 / max_producto)),
        }
        for nombre, cantidad in sorted(
            cantidades_productos.items(),
            key=lambda fila: (-fila[1], fila[0].casefold()),
        )[:10]
    ] if max_producto else []
    for indice, producto in enumerate(productos_mas_vendidos):
        producto['color'] = colores_grafico[indice % len(colores_grafico)]

    ventas_por_categoria = [
        {
            'nombre': nombre,
            'cantidad': cantidad,
            'porcentaje': max(2, round(cantidad * 100 / max_categoria)),
        }
        for nombre, cantidad in sorted(
            cantidades_categorias.items(),
            key=lambda fila: (-fila[1], fila[0].casefold()),
        )
    ] if max_categoria else []

    for indice, categoria in enumerate(ventas_por_categoria):
        categoria['color'] = colores_grafico[indice % len(colores_grafico)]
        categoria['participacion'] = round(categoria['cantidad'] * 100 / total_unidades) if total_unidades else 0

    max_pago = max(
        (resumen['total_cup'] for resumen in recaudacion_por_forma.values()),
        default=Decimal('0'),
    )
    pagos = [
        {
            **resumen,
            'porcentaje': max(2, round(resumen['total_cup'] * 100 / max_pago)),
            'participacion': round(resumen['total_cup'] * 100 / total_cup) if total_cup else 0,
            'color': colores_grafico[indice % len(colores_grafico)],
        }
        for indice, resumen in enumerate(recaudacion_por_forma.values())
    ] if max_pago else list(recaudacion_por_forma.values())

    pagos_con_recaudacion = [pago for pago in pagos if pago['total_cup'] > 0]
    segmentos_categorias = []
    acumulado_categorias = Decimal('0')
    for indice, categoria in enumerate(ventas_por_categoria):
        inicio = acumulado_categorias
        acumulado_categorias += Decimal(categoria['cantidad']) * Decimal('100') / Decimal(total_unidades)
        fin = Decimal('100') if indice == len(ventas_por_categoria) - 1 else acumulado_categorias
        segmentos_categorias.append(f"{categoria['color']} {inicio:.2f}% {fin:.2f}%")
    segmentos_pagos = []
    acumulado_pagos = Decimal('0')
    for indice, pago in enumerate(pagos_con_recaudacion):
        inicio = acumulado_pagos
        acumulado_pagos += pago['total_cup'] * Decimal('100') / total_cup
        fin = Decimal('100') if indice == len(pagos_con_recaudacion) - 1 else acumulado_pagos
        segmentos_pagos.append(f"{pago['color']} {inicio:.2f}% {fin:.2f}%")

    max_recaudacion = max(recaudacion_por_periodo.values(), default=Decimal('0'))
    recaudacion_grafico = [
        {
            'etiqueta': fecha.strftime('%b %Y' if (fecha_hasta - fecha_desde).days > 30 else '%d/%m'),
            'total': total,
            'porcentaje': max(2, round(total * 100 / max_recaudacion)),
        }
        for fecha, total in sorted(recaudacion_por_periodo.items())
    ] if max_recaudacion else []

    for indice, punto in enumerate(recaudacion_grafico):
        punto['color'] = colores_grafico[indice % len(colores_grafico)]

    context = {
        'periodo': periodo,
        'periodo_label': periodo_label,
        'fecha_desde': fecha_desde,
        'fecha_hasta': fecha_hasta,
        'agrupar_por_mes': not error_fechas and (fecha_hasta - fecha_desde).days > 30,
        'fecha_desde_raw': fecha_desde_raw,
        'fecha_hasta_raw': fecha_hasta_raw,
        'error_fechas': error_fechas,
        'total_cup': total_cup,
        'total_usd': total_usd,
        'total_efectivo_cup': total_efectivo_cup,
        'total_transferencia_cup': total_transferencia_cup,
        'cantidad_facturas': cantidad_facturas,
        'promedio_factura': total_cup / cantidad_facturas if cantidad_facturas else Decimal('0'),
        'total_unidades': total_unidades,
        'items_invalidos': items_invalidos,
        'productos_unicos': len(cantidades_productos),
        'productos_mas_vendidos': productos_mas_vendidos,
        'ventas_por_categoria': ventas_por_categoria,
        'pagos': pagos,
        'pagos_con_recaudacion': pagos_con_recaudacion,
        'categorias_donut': ', '.join(segmentos_categorias),
        'pagos_donut': ', '.join(segmentos_pagos),
        'recaudacion_grafico': recaudacion_grafico,
    }
    return render(request, 'usuarios/reportes.html', context)


@login_required
def historial_turno_detalle_view(request, turno_id):
    if request.user.role not in ROLES_CON_ACCESO_FINANCIERO:
        messages.error(request, 'Solo caja o el administrador pueden ver el detalle del turno actual.')
        return redirect('usuarios:dashboard')

    turno = get_object_or_404(Turno.objects.prefetch_related('usuarios'), pk=turno_id)
    if request.user.role == 'cajera' and turno.estado == 'cerrado':
        messages.error(request, 'La cajera no puede consultar turnos cerrados anteriores.')
        return redirect('usuarios:dashboard')
    data = turno.resumen_financiero
    datos_ipv = _datos_ipv_turno(turno)
    data['fondo_inicial_cup'] = turno.fondo_inicial_cup
    data['efectivo_esperado_cup'] = Decimal(datos_ipv['efectivo_final_cup'])
    context = {
        'turno': turno,
        'detalles': data,
        'resumen': turno.resumen or 'Sin resumen registrado.',
    }
    return render(request, 'usuarios/historial_turno_detalle.html', context)


@login_required
def ipv_turno(request, turno_id):
    """
    Devuelve el inventario perpetuo de ventas (IPV) de un turno concreto.
    Funciona tanto para turnos cerrados (histórico) como para el turno abierto (tiempo real).
    """
    if request.user.role not in ROLES_CON_ACCESO_FINANCIERO:
        return JsonResponse({'error': 'Sin permiso'}, status=403)

    turno = get_object_or_404(Turno, pk=turno_id)
    if request.user.role == 'cajera' and turno.estado == 'cerrado':
        return JsonResponse({'error': 'Sin permiso'}, status=403)
    return JsonResponse(_datos_ipv_turno(turno))


def _datos_ipv_turno(turno):
    cierre = turno.cierre if turno.estado == 'cerrado' and turno.cierre else timezone.now()
    registros = list(
        ProductoIPVTurno.objects.filter(turno=turno)
        .select_related('plato__categoria')
        .order_by('plato__nombre')
    )
    ventas = _ventas_facturadas_por_plato(turno, cierre, [registro.plato for registro in registros])
    facturas = Factura.objects.filter(creado_en__gte=turno.apertura, creado_en__lte=cierre)
    pagos = facturas.aggregate(
        efectivo=Sum('monto_efectivo_cup'),
        transferencia=Sum('monto_transferencia_cup'),
        usd=Sum('monto_usd'),
    )
    efectivo_ventas = pagos['efectivo'] or Decimal('0')
    fondo_inicial = turno.fondo_inicial_cup
    filas = []
    total_importe = Decimal('0')
    for registro in registros:
        plato = registro.plato
        venta = ventas[plato.pk]
        cantidad_vendida = venta['cantidad']
        importe_venta = venta['importe']
        saldo_final = (
            registro.cantidad_inicial
            + registro.entradas
            - cantidad_vendida
            - registro.merma_rotura
        )
        precio_venta = (
            importe_venta / cantidad_vendida
            if cantidad_vendida
            else plato.precio
        )
        total_importe += importe_venta
        filas.append({
            'id': registro.pk,
            'plato_id': plato.pk,
            'nombre': plato.nombre,
            'area': 'Carta',
            'categoria': plato.categoria.nombre if plato.categoria else 'Sin categoría',
            'unidad': 'Unidad',
            'saldo_inicial': str(registro.cantidad_inicial),
            'entradas': str(registro.entradas),
            'cantidad_existencia': str(saldo_final),
            'cantidad_vendida': str(cantidad_vendida),
            'saldo_final': str(saldo_final),
            'precio_venta': str(precio_venta),
            'importe_venta': str(importe_venta),
            'merma_rotura': str(registro.merma_rotura),
            'merma_url': reverse('usuarios:registrar_merma_ipv', args=[registro.pk]),
            'control_existencia': True,
        })

    return {
        'turno_id': turno.pk,
        'estado': turno.estado,
        'fecha': turno.fecha.strftime('%d/%m/%Y'),
        'actualizado_en': timezone.localtime(cierre).strftime('%H:%M:%S'),
        'total_importe': str(total_importe),
        'fondo_inicial_cup': f'{fondo_inicial:.2f}',
        'ventas_efectivo_cup': f'{efectivo_ventas:.2f}',
        'efectivo_final_cup': f'{fondo_inicial + efectivo_ventas:.2f}',
        'ventas_transferencia_cup': f'{pagos["transferencia"] or Decimal("0"):.2f}',
        'ventas_usd': f'{pagos["usd"] or Decimal("0"):.2f}',
        'filas': filas,
    }


IPV_COLUMNAS = (
    ('Descripción', 'nombre'),
    ('Categoría', 'categoria'),
    ('Saldo inicial', 'saldo_inicial'),
    ('Entradas', 'entradas'),
    ('En existencia', 'cantidad_existencia'),
    ('Saldo final', 'saldo_final'),
    ('Cant. vendida', 'cantidad_vendida'),
    ('Precio venta', 'precio_venta'),
    ('Importe', 'importe_venta'),
    ('Merma / rotura', 'merma_rotura'),
)


@login_required
def ipv_turno_pdf(request, turno_id):
    if request.user.role not in ROLES_CON_ACCESO_FINANCIERO:
        return JsonResponse({'error': 'Sin permiso'}, status=403)

    turno = get_object_or_404(Turno, pk=turno_id)
    if request.user.role == 'cajera' and turno.estado == 'cerrado':
        return JsonResponse({'error': 'Sin permiso'}, status=403)
    datos = _datos_ipv_turno(turno)
    buffer = BytesIO()
    pagina = landscape(A4)
    documento = SimpleDocTemplate(
        buffer,
        pagesize=pagina,
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
        title=f"IPV turno {turno.pk}",
    )
    estilos = getSampleStyleSheet()
    estilo_celda = ParagraphStyle(
        'IPVCelda',
        parent=estilos['BodyText'],
        fontName='Helvetica',
        fontSize=6,
        leading=7,
        spaceAfter=0,
    )
    estilo_numero = ParagraphStyle(
        'IPVNumero',
        parent=estilo_celda,
        alignment=TA_RIGHT,
    )
    estilo_encabezado = ParagraphStyle(
        'IPVEncabezado',
        parent=estilo_celda,
        fontName='Helvetica-Bold',
        textColor=colors.white,
    )
    encabezados = [Paragraph(escape(titulo), estilo_encabezado) for titulo, _ in IPV_COLUMNAS]
    filas_tabla = [encabezados]
    for fila in datos['filas']:
        filas_tabla.append([
            Paragraph(
                escape(str(fila[clave]) if fila[clave] is not None else '—'),
                estilo_celda if indice < 2 else estilo_numero,
            )
            for indice, (_, clave) in enumerate(IPV_COLUMNAS)
        ])
    filas_tabla.append([
        Paragraph('Total importe de ventas', estilo_celda),
        '', '', '', '', '', '', '',
        Paragraph(escape(datos['total_importe']), estilo_numero),
        '',
    ])
    ancho_disponible = pagina[0] - 20 * mm
    proporciones = (0.19, 0.12, 0.08, 0.08, 0.10, 0.08, 0.10, 0.09, 0.09, 0.07)
    tabla = Table(
        filas_tabla,
        colWidths=[ancho_disponible * proporcion for proporcion in proporciones],
        repeatRows=1,
    )
    estilos_tabla = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e293b')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#e2e8f0')),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('SPAN', (0, -1), (7, -1)),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#cbd5e1')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ALIGN', (2, 1), (-1, -1), 'RIGHT'),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]
    if datos['filas']:
        estilos_tabla.append(('BACKGROUND', (0, 1), (-1, -2), colors.HexColor('#f8fafc')))
        estilos_tabla.append(
            ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#f1f5f9')]),
        )
    tabla.setStyle(TableStyle(estilos_tabla))
    contenido = [
        Paragraph('Inventario Perpetuo de Ventas (IPV)', estilos['Title']),
        Paragraph(
            f"Turno {escape(datos['fecha'])} · "
            f"{'Abierto' if turno.estado == 'abierto' else 'Cerrado'} · "
            f"Actualizado {escape(datos['actualizado_en'])}",
            estilos['Normal'],
        ),
        Paragraph(
            f"Fondo inicial: {escape(datos['fondo_inicial_cup'])} CUP · "
            f"Efectivo vendido: {escape(datos['ventas_efectivo_cup'])} CUP · "
            f"Efectivo esperado: {escape(datos['efectivo_final_cup'])} CUP · "
            f"Transferencias: {escape(datos['ventas_transferencia_cup'])} CUP · "
            f"USD: {escape(datos['ventas_usd'])}",
            estilos['Normal'],
        ),
        Spacer(1, 8 * mm),
        tabla,
    ]
    documento.build(contenido)
    respuesta = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    respuesta['Content-Disposition'] = f'attachment; filename="ipv_turno_{turno.pk}.pdf"'
    return respuesta


@login_required
def ipv_turno_excel(request, turno_id):
    if request.user.role not in ROLES_CON_ACCESO_FINANCIERO:
        return JsonResponse({'error': 'Sin permiso'}, status=403)

    turno = get_object_or_404(Turno, pk=turno_id)
    if request.user.role == 'cajera' and turno.estado == 'cerrado':
        return JsonResponse({'error': 'Sin permiso'}, status=403)
    datos = _datos_ipv_turno(turno)
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'IPV'
    hoja.append(['Inventario Perpetuo de Ventas (IPV)'])
    hoja.append([
        'Turno',
        datos['fecha'],
        'Estado',
        'Abierto' if turno.estado == 'abierto' else 'Cerrado',
        'Actualizado',
        datos['actualizado_en'],
    ])
    hoja.append([
        'Fondo inicial CUP', Decimal(datos['fondo_inicial_cup']),
        'Efectivo vendido CUP', Decimal(datos['ventas_efectivo_cup']),
        'Efectivo esperado CUP', Decimal(datos['efectivo_final_cup']),
        'Transferencias CUP', Decimal(datos['ventas_transferencia_cup']),
        'Ventas USD', Decimal(datos['ventas_usd']),
    ])
    hoja.append([titulo for titulo, _ in IPV_COLUMNAS])
    for celda in hoja[4]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor='1E293B')
        celda.alignment = Alignment(horizontal='center', wrap_text=True)

    for fila in datos['filas']:
        valores = []
        for _, clave in IPV_COLUMNAS:
            valor = fila[clave]
            if valor is None:
                valores.append(None)
            elif clave in {
                'saldo_inicial',
                'entradas',
                'cantidad_existencia',
                'saldo_final',
                'cantidad_vendida',
                'precio_venta',
                'importe_venta',
                'merma_rotura',
            }:
                valores.append(Decimal(valor))
            else:
                valores.append(str(valor))
        hoja.append(valores)

    fila_total = hoja.max_row + 1
    hoja.cell(row=fila_total, column=1, value='Total importe de ventas')
    hoja.cell(row=fila_total, column=9, value=Decimal(datos['total_importe']))
    for celda in hoja[fila_total]:
        celda.font = Font(bold=True)
        celda.fill = PatternFill('solid', fgColor='E2E8F0')

    hoja.freeze_panes = 'A5'
    hoja.auto_filter.ref = f'A4:J{fila_total - 1}'
    anchos = (32, 22, 16, 14, 16, 14, 16, 16, 16, 18)
    for indice, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[chr(64 + indice)].width = ancho
    for fila in hoja.iter_rows(min_row=5, max_row=fila_total):
        for indice in range(3, 11):
            celda = fila[indice - 1]
            if celda.value is not None:
                celda.number_format = '#,##0.00'
        for indice in (1, 2):
            celda = fila[indice - 1]
            if celda.value is not None:
                celda.data_type = 's'

    buffer = BytesIO()
    libro.save(buffer)
    respuesta = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    respuesta['Content-Disposition'] = f'attachment; filename="ipv_turno_{turno.pk}.xlsx"'
    return respuesta
